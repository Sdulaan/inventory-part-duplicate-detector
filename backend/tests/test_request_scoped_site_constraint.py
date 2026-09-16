"""Request-scoped CONTRACT equality without durable identity contamination."""

from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from app.core.config import Settings
from app.db.models import (
    IdentityEvidenceEdgeSnapshot,
    IdentityNeighborProposal,
    IdentityResolutionConstraintInput,
)
from app.engine.identity_edge import IdentityEdgeClass
from app.resolution.contracts import (
    IdentityResolutionEvidenceEdge,
    IdentityResolutionInput,
    IdentityResolutionNeighborhood,
    ResolverConfiguration,
)
from app.resolution.request_constraints import (
    CONTRACT_GROUP_CONSTRAINT,
    contract_pair_is_compatible,
)
from app.resolution.resolver import resolve_identity_groups
from app.services.canonical_record_service import CanonicalScanRecord
from app.services.identity_read_service import IdentityReadService
from app.services.identity_read_xlsx_export_service import (
    authority_selected_system_groups_to_xlsx,
)
from app.services.scan_runner import ScanRunner


def configuration():
    return Settings(
        llm_provider="none",
        llm_demo_enabled=False,
        identity_orchestration_mode="group_first_primary",
        group_first_shadow_comparison_enabled=False,
        hybrid_retrieval_enabled=True,
        local_embedding_enabled=False,
        hybrid_retrieval_lexical_top_k=10,
        hybrid_retrieval_vector_top_k=10,
        hybrid_retrieval_final_top_k=10,
        hybrid_retrieval_max_pairs_per_scan=100,
        hybrid_retrieval_family_max=30,
        hybrid_retrieval_tier_a_max=100,
        hybrid_retrieval_tier_b_max=100,
        hybrid_retrieval_tier_c_max=100,
        identity_neighborhood_max_members=20,
    )


def two_site_records():
    return pd.DataFrame([
        {
            "PART_NO": "A",
            "DESCRIPTION": "SKF 6205 BEARING",
            "CONTRACT": "ST001",
            "UNIT_MEAS": "PCS",
        },
        {
            "PART_NO": "B",
            "DESCRIPTION": "SKF BEARING 6205",
            "CONTRACT": "ST002",
            "UNIT_MEAS": "PCS",
        },
    ])


def run(db, selected_fields, name):
    return ScanRunner(db, configuration()).run(
        two_site_records(),
        name,
        selected_fields,
        75,
        scan_mode="SAME_SITE_DUPLICATE",
        orchestration_mode="group_first_primary",
    )[0]


def canonical(record_id, contract):
    return CanonicalScanRecord(
        record_id=record_id,
        scan_id=1,
        source_row_index=record_id - 1,
        record_ref_key=f"record-{record_id}",
        source_record_fingerprint=f"source-{record_id}",
        part_no=chr(64 + record_id),
        description="SKF 6205 BEARING",
        contract=contract,
        uom="PCS",
        type_code=None,
        prime_commodity=None,
        second_commodity=None,
        accounting_group=None,
        part_product_code=None,
        part_product_family=None,
        product_category_id=None,
        hsn_sac_code=None,
        hazard_code=None,
        normalized_part_no=chr(64 + record_id),
        normalized_description="SKF 6205 BEARING",
        normalization_version="site-constraint-test-v1",
    )


def resolver_input(constraints):
    records = (canonical(1, "ST001"), canonical(2, "ST002"))
    return IdentityResolutionInput(
        scan_id=1,
        discovery_run_id=1,
        evidence_run_id=1,
        canonical_records=records,
        identity_neighborhoods=(
            IdentityResolutionNeighborhood("n-1", 1, 1, (1, 2), False, False),
        ),
        machine_evidence_edges=(
            IdentityResolutionEvidenceEdge(
                scan_id=1,
                evidence_run_id=1,
                record_id_1=1,
                record_id_2=2,
                edge_class=IdentityEdgeClass.STRONG_SUPPORT,
                reason_codes=("TEST_STRONG",),
                evidence_fingerprint="edge-1-2",
                generic_only=False,
            ),
        ),
        human_constraints=(),
        resolver_algorithm_version="site-constraint-test-v1",
        resolver_configuration=ResolverConfiguration(20, 40, 8, "test-v1"),
        request_scoped_group_constraints=constraints,
    )


def test_missing_contract_handling_is_conservative_and_case_insensitive():
    assert contract_pair_is_compatible({"CONTRACT": " ST001 "}, {"CONTRACT": "st001"})
    assert not contract_pair_is_compatible({"CONTRACT": "ST001"}, {"CONTRACT": "ST002"})
    assert not contract_pair_is_compatible({"CONTRACT": "ST001"}, {"CONTRACT": ""})
    assert contract_pair_is_compatible({"CONTRACT": None}, {"CONTRACT": "  "})


def test_gf5_rejects_cross_site_partition_only_when_request_constraint_is_active():
    unrestricted = resolve_identity_groups(resolver_input(()), None)
    restricted = resolve_identity_groups(
        resolver_input((CONTRACT_GROUP_CONSTRAINT,)), None
    )

    assert [group.member_record_ids for group in unrestricted.accepted_groups] == [(1, 2)]
    assert restricted.accepted_groups == ()
    assert restricted.unassigned_record_ids == (1, 2)
    assert restricted.conflicts == ()


def test_sequential_scans_prove_early_pruning_and_no_persistent_site_constraint(
    db, monkeypatch,
):
    def provider_called(*_args, **_kwargs):
        raise AssertionError("site constraint test invoked a provider")

    monkeypatch.setattr("app.llm.factory.create_llm_provider", provider_called)
    monkeypatch.setattr(
        "app.llm.groq_group_provider.create_group_advisory_provider", provider_called
    )

    selected = run(db, ["CONTRACT", "UNIT_MEAS"], "site selected")
    unselected = run(db, ["UNIT_MEAS"], "site unselected")

    selected_snapshot = IdentityReadService(db).load_identity_read_snapshot(selected.id)
    unselected_snapshot = IdentityReadService(db).load_identity_read_snapshot(unselected.id)
    assert selected_snapshot.groups == ()
    assert len(unselected_snapshot.groups) == 1
    assert {member.contract for member in unselected_snapshot.groups[0].members} == {
        "ST001", "ST002",
    }

    assert db.query(IdentityNeighborProposal).filter_by(scan_id=selected.id).count() == 0
    assert db.query(IdentityEvidenceEdgeSnapshot).filter_by(scan_id=selected.id).count() == 0
    assert db.query(IdentityNeighborProposal).filter_by(scan_id=unselected.id).count() == 1
    assert db.query(IdentityResolutionConstraintInput).count() == 0


def test_selected_site_xlsx_contains_only_single_site_groups(db):
    rows = pd.DataFrame([
        {
            "PART_NO": "A",
            "DESCRIPTION": "SKF 6205 BEARING",
            "CONTRACT": "ST001",
            "UNIT_MEAS": "PCS",
        },
        {
            "PART_NO": "B",
            "DESCRIPTION": "SKF BEARING 6205",
            "CONTRACT": "ST001",
            "UNIT_MEAS": "PCS",
        },
        {
            "PART_NO": "C",
            "DESCRIPTION": "SKF BEARING 6205",
            "CONTRACT": "ST002",
            "UNIT_MEAS": "PCS",
        },
    ])
    scan = ScanRunner(db, configuration()).run(
        rows, "site XLSX", ["CONTRACT", "UNIT_MEAS"], 75,
        scan_mode="SAME_SITE_DUPLICATE",
        orchestration_mode="group_first_primary",
    )[0]
    snapshot = IdentityReadService(db).load_identity_read_snapshot(scan.id)
    assert snapshot.groups
    assert all(
        len({member.contract.casefold() for member in group.members if member.contract}) <= 1
        for group in snapshot.groups
    )

    workbook = load_workbook(BytesIO(authority_selected_system_groups_to_xlsx(db, scan.id)))
    assert workbook.sheetnames == [
        "Overview", "Review Groups", "Group Index", "Detailed Data",
        "Technical Reference",
    ]
    sheet = workbook["Detailed Data"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    sites_by_group = {}
    for row in range(2, sheet.max_row + 1):
        group = sheet.cell(row, headers["Group"]).value
        site = sheet.cell(row, headers["Site"]).value
        sites_by_group.setdefault(group, set()).add(str(site).strip().casefold())
    assert sites_by_group
    assert all(len(sites) <= 1 for sites in sites_by_group.values())
