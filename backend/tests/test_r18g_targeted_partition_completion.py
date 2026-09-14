import csv

import pytest
from openpyxl import Workbook, load_workbook

from app.benchmarks.r18g_development_group_evaluation import (
    REFERENCE_FIELDS,
    REFERENCE_TYPE,
    REVIEW_HEADERS,
)
from app.benchmarks.r18g_group_human_evidence import PARTITION_IDS, sha256_file
from app.benchmarks.r18g_targeted_partition_completion import (
    INTERNAL_MAPPING_FIELDS,
    MEMBER_HEADERS,
    TARGET_IDS,
    TargetedPartitionPackageError,
    comparison_set_id,
    load_target_sources,
    prepare_targeted_package,
)


def _write_csv(path, fields, rows):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _sources(tmp_path):
    submission = tmp_path / "development_completed.xlsx"
    reference = tmp_path / "development_reference.csv"
    mapping = tmp_path / "internal_mapping.csv"
    workbook = Workbook()
    review = workbook.active
    review.title = "Group Review"
    review.append(REVIEW_HEADERS)
    members = workbook.create_sheet("Members")
    members.append(MEMBER_HEADERS)
    reference_rows = []
    mapping_rows = []
    for group_index, group_id in enumerate(TARGET_IDS):
        count = 5 if group_index == 0 else 4
        raw = "MIXED_GROUP_MULTIPLE_IDENTITIES" if group_index == 0 else "NOT_ONE_IDENTITY"
        normalized = "NOT_ONE_IDENTITY"
        confidence = "MEDIUM"
        reason = "MIXED_GROUP_MULTIPLE_IDENTITIES"
        comment = f"Prior immutable judgment for {group_id}."
        basis = "MIXED"
        review.append([group_id, count, raw, confidence, reason, comment, basis])
        refs = []
        for member_no in range(1, count + 1):
            refs.append(f"REF-{group_index}-{member_no}")
            members.append([
                group_id, member_no, f"PART-{group_index}-{member_no}", "",
                f"Description {member_no}", "", "", "", "EA", "STOCK",
                "SITE-A",
            ])
        reference_rows.append({
            "review_group_id": group_id, "member_count": count,
            "raw_group_label": raw, "normalized_group_label": normalized,
            "confidence": confidence, "reason_code": reason,
            "reviewer_comment": comment, "comment_basis": basis,
            "partition_detail_status": "UNAVAILABLE",
            "reference_type": REFERENCE_TYPE, "submission_sha256": "f" * 64,
            "dataset_version": "r18g-group-human-evidence-v1",
            "normalization_reason": "",
        })
        mapping_rows.append({
            "review_group_id": group_id,
            "hypothesis_fingerprint": f"fingerprint-{group_index}",
            "source_kind": "ADR_PARTITION_CHALLENGER",
            "current_or_shadow": "SHADOW_ADR_ORDER",
            "current_group_id": "", "work_unit_id": f"work-{group_index}",
            "system_status": "", "group_size": count, "site_count": 1,
            "evidence_provenance": "MIXED_EVIDENCE", "strong_edge_count": 1,
            "native_review_edge_count": 1,
            "r18c_demoted_review_edge_count": 0, "cannot_link_count": 0,
            "generic_only_state": "false", "partition_sensitive": "true",
            "deferred_context": "false", "conflict_context": "false",
            "sampling_stratum": "TEST", "development_or_holdout": "DEVELOPMENT",
            "source_member_refs": "|".join(refs),
        })
    workbook.save(submission)
    _write_csv(reference, REFERENCE_FIELDS, reference_rows)
    mapping_fields = tuple(mapping_rows[0].keys())
    _write_csv(mapping, mapping_fields, mapping_rows)
    return submission, reference, mapping


def test_exact_two_group_source_selection_and_judgment_preservation(tmp_path):
    submission, reference, mapping = _sources(tmp_path)
    references, mappings, members = load_target_sources(
        submission, reference, mapping
    )
    assert tuple(references) == TARGET_IDS
    assert tuple(mappings) == TARGET_IDS
    assert len(members) == 9
    assert sum(row[0] == TARGET_IDS[0] for row in members) == 5
    assert sum(row[0] == TARGET_IDS[1] for row in members) == 4
    assert references[TARGET_IDS[0]]["raw_group_label"] == (
        "MIXED_GROUP_MULTIPLE_IDENTITIES"
    )


def test_targeted_workbook_is_blinded_empty_protected_and_deterministic(tmp_path):
    submission, reference, mapping = _sources(tmp_path)
    output = tmp_path / "output"
    first = prepare_targeted_package(submission, reference, mapping, output)
    workbook_path = output / "r18g_targeted_partition_completion.xlsx"
    mapping_path = output / "r18g_targeted_partition_internal_mapping.csv"
    first_hashes = (sha256_file(workbook_path), sha256_file(mapping_path))
    second = prepare_targeted_package(submission, reference, mapping, output)
    assert (sha256_file(workbook_path), sha256_file(mapping_path)) == first_hashes
    assert first["validation"] == second["validation"]
    assert first["validation"]["target_ids"] == list(TARGET_IDS)
    assert first["validation"]["prefilled_partition_ids"] == 0
    assert first["validation"]["detector_fields_visible"] == 0

    workbook = load_workbook(workbook_path, data_only=False)
    assert workbook.sheetnames == ["Instructions", "Groups", "Members", "Partition"]
    assert all(workbook[name].protection.sheet for name in workbook.sheetnames)
    partition = workbook["Partition"]
    assert all(
        cell.value in (None, "") and cell.data_type != "f" and not cell.protection.locked
        for row in partition.iter_rows(min_row=2, min_col=4, max_col=5)
        for cell in row
    )
    assert {
        item.formula1 for item in partition.data_validations.dataValidation
    } == {'"' + ",".join(PARTITION_IDS) + '"'}
    visible = "\n".join(
        str(cell.value or "") for sheet in workbook for row in sheet.iter_rows()
        for cell in row
    ).casefold()
    assert "holdout" not in visible
    assert "current_or_shadow" not in visible
    assert "work_unit_id" not in visible

    with mapping_path.open(encoding="utf-8", newline="") as handle:
        internal = list(csv.DictReader(handle))
    assert tuple(internal[0]) == INTERNAL_MAPPING_FIELDS
    assert [row["review_group_id"] for row in internal] == list(TARGET_IDS)


def test_comparison_identity_is_stable():
    assert comparison_set_id("WORK", "A|B") == comparison_set_id("WORK", "A|B")
    assert comparison_set_id("WORK", "A|B") != comparison_set_id("WORK", "A|C")


def test_member_count_or_prior_judgment_mismatch_fails_closed(tmp_path):
    submission, reference, mapping = _sources(tmp_path)
    workbook = load_workbook(submission)
    workbook["Group Review"]["F2"] = "Changed judgment"
    workbook.save(submission)
    with pytest.raises(TargetedPartitionPackageError, match="prior judgment mismatch"):
        load_target_sources(submission, reference, mapping)


def test_holdout_named_output_is_rejected(tmp_path):
    submission, reference, mapping = _sources(tmp_path)
    with pytest.raises(ValueError, match="R18G_B_HOLDOUT_FIREWALL_BLOCKED"):
        prepare_targeted_package(
            submission, reference, mapping, tmp_path / "sealed_holdout_output"
        )
