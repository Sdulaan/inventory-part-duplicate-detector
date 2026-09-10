"""Focused tests for consumed R18G-C holdout validation."""

from __future__ import annotations

import csv

import pytest
from openpyxl import Workbook, load_workbook

from app.benchmarks.r18g_group_human_evidence import VISIBLE_FIELDS
from app.benchmarks.r18g_holdout_validation import (
    HOLDOUT_STATUS,
    MIXED_LABEL,
    MIXED_NORMALIZATION_REASON,
    HoldoutLabelVocabularyError,
    HoldoutReferenceError,
    canonicalize_partition,
    classify_human_group,
    classify_severity,
    graduation_decision,
    load_holdout_mapping,
    normalize_label,
    validate_and_recover,
)


def _workbook(path, *, completed):
    workbook = Workbook()
    instructions = workbook.active
    instructions.title = "Instructions"
    instructions.append(["R18G blinded whole-group review"])
    instructions.append(["dataset_version", "r18g-group-human-evidence-v1"])
    review = workbook.create_sheet("Group Review")
    review.append((
        "review_group_id", "member_count", "group_label", "confidence",
        "reason_code", "reviewer_comment", "comment_basis",
    ))
    members = workbook.create_sheet("Members")
    members.append(("review_group_id", "member_no", *VISIBLE_FIELDS))
    partition = workbook.create_sheet("Partition")
    partition.append((
        "review_group_id", "member_no", "part_number", "human_partition_id",
        "partition_comment",
    ))
    for index in range(16):
        group_id = f"R18G-{index:016X}"
        human = (
            "SAME_ONE_IDENTITY", "HIGH", "SAME_ALIAS_OR_NAMING_VARIANT",
            "Same source-visible identity.", "SOURCE_VISIBLE",
        ) if completed else ("", "", "", "", "")
        review.append((group_id, 2, *human))
        for member_no in (1, 2):
            visible = [f"PN-{index}-{member_no}"] + [""] * (len(VISIBLE_FIELDS) - 1)
            members.append((group_id, member_no, *visible))
            partition.append((group_id, member_no, visible[0], "", ""))
    workbook.save(path)


def _mapping(path):
    fields = (
        "review_group_id", "development_or_holdout", "hypothesis_fingerprint",
        "source_kind", "current_or_shadow", "current_group_id", "work_unit_id",
        "system_status", "group_size", "site_count", "evidence_provenance",
        "strong_edge_count", "native_review_edge_count",
        "r18c_demoted_review_edge_count", "cannot_link_count",
        "generic_only_state", "partition_sensitive", "deferred_context",
        "conflict_context", "sampling_stratum", "source_member_refs",
    )
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerow({
            "review_group_id": "R18G-DEVELOPMENT", "development_or_holdout": "DEVELOPMENT",
        })
        for index in range(16):
            writer.writerow({
                "review_group_id": f"R18G-{index:016X}",
                "development_or_holdout": "SEALED_HOLDOUT",
                "hypothesis_fingerprint": f"fp-{index}",
                "source_kind": "ACCEPTED_LIKELY_GROUP",
                "current_or_shadow": "CURRENT",
                "current_group_id": f"group-{index}",
                "work_unit_id": f"unit-{index}",
                "system_status": "LIKELY_DUPLICATE_GROUP",
                "group_size": "2", "site_count": "1",
                "evidence_provenance": "TEST", "strong_edge_count": "1",
                "native_review_edge_count": "0",
                "r18c_demoted_review_edge_count": "0", "cannot_link_count": "0",
                "generic_only_state": "false", "partition_sensitive": "false",
                "deferred_context": "false", "conflict_context": "false",
                "sampling_stratum": "TEST",
                "source_member_refs": f"REF-{index}-1|REF-{index}-2",
            })


def _joined(*, likely_not_one=0, review_not_one=0, cannot_link=0, duplicate=False):
    rows = []
    for index in range(4):
        likely = index < 2
        not_one = (
            index < likely_not_one if likely
            else index - 2 < review_not_one
        )
        refs = "A|B" if duplicate and index < 2 else f"R{index}A|R{index}B"
        rows.append({
            "source_kind": "ACCEPTED_LIKELY_GROUP" if likely else "ACCEPTED_REVIEW_GROUP",
            "system_status": "LIKELY_DUPLICATE_GROUP" if likely else "POSSIBLE_DUPLICATE_GROUP_REVIEW",
            "normalized_group_label": "NOT_ONE_IDENTITY" if not_one else "SAME_ONE_IDENTITY",
            "confidence": "MEDIUM",
            "cannot_link_count": str(cannot_link if index == 0 else 0),
            "source_member_refs": refs,
            "partition_detail_status": "FORCED_BINARY_SEPARATION" if not_one else "ONE_IDENTITY_IMPLIED_BY_GROUP_LABEL",
            "partition_relation": "NOT_COMPARABLE",
        })
    return rows


def test_exact_16_id_reconciliation_raw_preservation_and_disjoint_mapping(tmp_path):
    blank = tmp_path / "blank.xlsx"
    submission = tmp_path / "completed.xlsx"
    mapping = tmp_path / "mapping.csv"
    _workbook(blank, completed=False)
    _workbook(submission, completed=True)
    _mapping(mapping)
    reference, partitions, integrity = validate_and_recover(submission, blank)
    assert len(reference) == 16
    assert len(partitions) == 32
    assert integrity["unique_id_count"] == 16
    assert reference[0]["raw_group_label"] == "SAME_ONE_IDENTITY"
    rows = load_holdout_mapping(mapping, (row["review_group_id"] for row in reference))
    assert len(rows) == 16
    assert "R18G-DEVELOPMENT" not in rows


def test_missing_or_foreign_holdout_id_fails_closed(tmp_path):
    blank = tmp_path / "blank.xlsx"
    submission = tmp_path / "completed.xlsx"
    _workbook(blank, completed=False)
    _workbook(submission, completed=True)
    workbook = load_workbook(submission)
    workbook["Group Review"]["A2"] = "FOREIGN"
    workbook.save(submission)
    with pytest.raises(HoldoutReferenceError, match="missing or foreign"):
        validate_and_recover(submission, blank)


def test_strict_mixed_token_recovery_preserves_raw_semantics():
    normalized, reason = normalize_label(
        MIXED_LABEL,
        reason_code=MIXED_LABEL,
        member_count=3,
        reviewer_comment="Multiple identities are present.",
    )
    assert normalized == "NOT_ONE_IDENTITY"
    assert reason == MIXED_NORMALIZATION_REASON
    with pytest.raises(HoldoutLabelVocabularyError):
        normalize_label(
            MIXED_LABEL,
            reason_code="OTHER",
            member_count=3,
            reviewer_comment="Multiple identities are present.",
        )


def test_partition_canonicalization_uses_member_set_equivalence():
    refs = ("B", "A", "D", "C")
    assert canonicalize_partition(refs, ("P9", "P9", "P4", "P4")) == (
        "P1", "P1", "P2", "P2"
    )
    assert canonicalize_partition(refs, ("", "", "", "")) == ("", "", "", "")


@pytest.mark.parametrize(
    ("label", "expected"),
    (
        ("SAME_ONE_IDENTITY", "HUMAN_SUPPORTED_GROUP"),
        ("NOT_ONE_IDENTITY", "HUMAN_REJECTED_GROUP"),
        ("INSUFFICIENT_INFORMATION", "HUMAN_UNRESOLVED_GROUP"),
    ),
)
def test_accepted_group_human_classification(label, expected):
    assert classify_human_group(label) == expected


def test_likely_and_review_severity_remain_distinct():
    assert classify_severity(
        "LIKELY_DUPLICATE_GROUP", "NOT_ONE_IDENTITY", "HIGH"
    ) == "CRITICAL_GROUP_QUALITY_FAILURE"
    assert classify_severity(
        "LIKELY_DUPLICATE_GROUP", "NOT_ONE_IDENTITY", "MEDIUM"
    ) == "SERIOUS_GROUP_QUALITY_FAILURE"
    assert classify_severity(
        "POSSIBLE_DUPLICATE_GROUP_REVIEW", "NOT_ONE_IDENTITY", "HIGH"
    ) == "REVIEW_BURDEN_FALSE_HYPOTHESIS"


def test_cannot_link_and_duplicate_membership_fail_graduation():
    classification, _, gates, facts = graduation_decision(
        _joined(cannot_link=1, duplicate=True)
    )
    assert classification == "R18G_C_HOLDOUT_REVEALS_CRITICAL_GROUP_QUALITY_DEFECT"
    assert gates["A_no_accepted_cannot_link"] is False
    assert gates["B_no_duplicate_accepted_membership"] is False
    assert facts["accepted_cannot_link_violations"] == 1


def test_graduation_gate_is_deterministic_for_repeated_likely_rejections():
    first = graduation_decision(_joined(likely_not_one=2, review_not_one=1))
    second = graduation_decision(_joined(likely_not_one=2, review_not_one=1))
    assert first == second
    assert first[0] == "R18G_C_HOLDOUT_REVEALS_CRITICAL_GROUP_QUALITY_DEFECT"
    assert first[1] == "ARCHITECT_REVIEW_HOLDOUT_DEFECT"
    assert first[2]["E_no_generalized_defect_invalidating_policy"] is False


def test_review_noise_without_likely_failure_is_acceptable_with_caution():
    classification, next_step, gates, _ = graduation_decision(
        _joined(review_not_one=1)
    )
    assert all(gates.values())
    assert classification == "R18G_C_HOLDOUT_ACCEPTABLE_WITH_REVIEW_CAUTION"
    assert next_step == "R19A_GROUP_CONFIDENCE_DESIGN"


def test_holdout_consumed_governance_status_is_frozen():
    assert HOLDOUT_STATUS == "CONSUMED_FOR_GF5_GROUP_VALIDATION"
