import csv
import json

import pytest
from openpyxl import load_workbook

from app.benchmarks.r18g_development_group_evaluation import (
    DevelopmentReferenceError,
    LabelVocabularyError,
    assert_holdout_firewall,
    evaluate_development,
    normalize_label,
    partition_detail_status,
    partition_policy_decision,
    validate_and_recover,
)
from app.benchmarks.r18g_group_human_evidence import (
    GroupHypothesis,
    Member,
    fingerprint,
    review_group_id,
    write_manifest,
    write_mapping,
    write_workbook,
)


def _member(number):
    return Member(
        number, f"REF-{number:03d}", number, f"PART-{number}", "", f"Item {number}",
        "", "", "", "EA", "STOCK", "SITE-A",
    )


def _group(number, *, source="ACCEPTED_REVIEW_GROUP", members=None,
           shadow=False, partition=False):
    values = members or (_member(number * 10), _member(number * 10 + 1))
    return GroupHypothesis(
        fingerprint((item.stable_ref for item in values), source), tuple(values), source,
        "SHADOW_ADR_ORDER" if shadow else "CURRENT", f"CURRENT-{number}",
        f"UNIT-{number}", "POSSIBLE_DUPLICATE_GROUP_REVIEW", "MIXED_EVIDENCE",
        0, 1, 0, 0, False, partition, source == "DEFERRED_FAMILY_CANDIDATE",
        source == "CONFLICT_CONTEXT_CANDIDATE",
    )


def _package(tmp_path):
    shared = (_member(900), _member(901), _member(902))
    shadow = _group(
        90, source="ADR_PARTITION_CHALLENGER", members=shared,
        shadow=True, partition=True,
    )
    current = _group(91, source="DEFERRED_FAMILY_CANDIDATE", members=shared)
    bicycle = _group(
        92, source="GENERIC_BICYCLE_FAMILY",
        members=tuple(_member(920 + index) for index in range(7)),
    )
    groups = [shadow, current, bicycle]
    groups.extend(_group(index) for index in range(45))
    groups = tuple(groups)
    blank = tmp_path / "development_blank.xlsx"
    submission = tmp_path / "development_completed.xlsx"
    mapping = tmp_path / "mapping.csv"
    manifest = tmp_path / "manifest.json"
    write_workbook(blank, groups, "DEVELOPMENT")
    submission.write_bytes(blank.read_bytes())
    write_mapping(mapping, groups, ())
    write_manifest(manifest, {"dataset_version": "test"})
    workbook = load_workbook(submission)
    review = workbook["Group Review"]
    for row in review.iter_rows(min_row=2):
        row[2].value = "SAME_ONE_IDENTITY"
        row[3].value = "HIGH"
        row[4].value = "SAME_ALIAS_OR_NAMING_VARIANT"
        row[5].value = "Visible aliases refer to one item."
        row[6].value = "SOURCE_VISIBLE"
        if row[0].value == review_group_id(shadow):
            row[2].value = "MIXED_GROUP_MULTIPLE_IDENTITIES"
            row[3].value = "MEDIUM"
            row[4].value = "MIXED_GROUP_MULTIPLE_IDENTITIES"
            row[5].value = "The group visibly contains multiple identities."
            row[6].value = "MIXED"
    workbook.save(submission)
    return submission, blank, mapping, manifest, shadow, bicycle


def test_strict_mixed_label_recovery_preserves_raw_semantics():
    normalized, reason = normalize_label(
        "MIXED_GROUP_MULTIPLE_IDENTITIES",
        reason_code="MIXED_GROUP_MULTIPLE_IDENTITIES", member_count=3,
        reviewer_comment="Multiple identities are present.", comment_basis="MIXED",
    )
    assert normalized == "NOT_ONE_IDENTITY"
    assert reason == "HUMAN_USED_REASON_TOKEN_AS_SEMANTIC_NOT_ONE_IDENTITY_LABEL"
    with pytest.raises(LabelVocabularyError):
        normalize_label(
            "MIXED_GROUP_MULTIPLE_IDENTITIES",
            reason_code="OTHER", member_count=3, reviewer_comment="Multiple.",
            comment_basis="MIXED",
        )
    with pytest.raises(LabelVocabularyError):
        normalize_label(
            "SOMETHING_ELSE", reason_code="OTHER", member_count=2,
            reviewer_comment="Text", comment_basis="UNCLEAR",
        )


def test_partition_detail_derivation_does_not_infer_multi_member_split():
    assert partition_detail_status("SAME_ONE_IDENTITY", 4, (None,) * 4) == (
        "ONE_IDENTITY_IMPLIED_BY_GROUP_LABEL"
    )
    assert partition_detail_status("NOT_ONE_IDENTITY", 2, (None, None)) == (
        "FORCED_BINARY_SEPARATION"
    )
    assert partition_detail_status("NOT_ONE_IDENTITY", 4, (None,) * 4) == "UNAVAILABLE"
    assert partition_detail_status("INSUFFICIENT_INFORMATION", 3, (None,) * 3) == "UNRESOLVED"
    with pytest.raises(DevelopmentReferenceError, match="partition vocabulary"):
        partition_detail_status("NOT_ONE_IDENTITY", 2, ("GROUP-A", "GROUP-B"))


def test_holdout_firewall_rejects_any_holdout_named_input(tmp_path):
    with pytest.raises(ValueError, match="R18G_B_HOLDOUT_FIREWALL_BLOCKED"):
        assert_holdout_firewall(tmp_path / "sealed_holdout.xlsx")


def test_partition_decision_requires_only_targeted_missing_detail():
    assert partition_policy_decision((), ({"review_group_id": "R18G-X"},)) == (
        "R18G_B_PARTITION_DETAIL_REQUIRED",
        "R18G_B1_TARGETED_PARTITION_COMPLETION",
    )


def test_development_integrity_48_ids_raw_preservation_and_mapping_join(tmp_path):
    submission, blank, mapping, manifest, shadow, bicycle = _package(tmp_path)
    reference, integrity = validate_and_recover(submission, blank)
    assert len(reference) == 48
    assert integrity["unique_id_count"] == 48
    recovered = {row["review_group_id"]: row for row in reference}
    row = recovered[review_group_id(shadow)]
    assert row["raw_group_label"] == "MIXED_GROUP_MULTIPLE_IDENTITIES"
    assert row["normalized_group_label"] == "NOT_ONE_IDENTITY"
    assert row["partition_detail_status"] == "UNAVAILABLE"

    result = evaluate_development(
        submission, blank, mapping, manifest, tmp_path / "out"
    )
    assert result["integrity"]["missing_id_count"] == 0
    assert result["mixed_label_recovery_count"] == 1
    assert result["partition_comparison_count"] == 1
    assert result["partition_comparison_outcomes"] == {
        "CURRENT_ONLY_HUMAN_SUPPORTED": 1
    }
    assert result["targeted_partition_followup"] == [{
        "review_group_id": review_group_id(shadow), "member_count": 3,
    }]
    assert result["bicycle_finding"]["review_group_id"] == review_group_id(bicycle)
    assert result["comment_basis_distribution"]["MIXED"] == 1
    assert result["holdout_firewall"] == "SEALED_HUMAN_LABELS_UNSEEN"

    with (tmp_path / "out" / "r18g_development_group_evaluation.csv").open(
        encoding="utf-8", newline=""
    ) as handle:
        rows = list(csv.DictReader(handle))
    assert len(rows) == 48
    by_id = {row["review_group_id"]: row for row in rows}
    assert by_id[review_group_id(shadow)]["current_or_shadow"] == "SHADOW_ADR_ORDER"
    assert by_id[review_group_id(bicycle)]["source_kind"] == "GENERIC_BICYCLE_FAMILY"
    assert json.loads(
        (tmp_path / "out" / "r18g_development_summary.json").read_text(encoding="utf-8")
    )["provider_calls"] == 0


def test_source_member_tampering_fails_closed(tmp_path):
    submission, blank, *_ = _package(tmp_path)
    workbook = load_workbook(submission)
    workbook["Members"]["C2"] = "CHANGED"
    workbook.save(submission)
    with pytest.raises(DevelopmentReferenceError, match="source members changed"):
        validate_and_recover(submission, blank)


def test_foreign_or_missing_review_id_fails_closed(tmp_path):
    submission, blank, *_ = _package(tmp_path)
    workbook = load_workbook(submission)
    workbook["Group Review"]["A2"] = "FOREIGN-ID"
    workbook.save(submission)
    with pytest.raises(DevelopmentReferenceError, match="48-ID reconciliation"):
        validate_and_recover(submission, blank)
