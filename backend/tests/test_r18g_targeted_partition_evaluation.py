import csv

import pytest
from openpyxl import Workbook, load_workbook

from app.benchmarks.r18g_development_group_evaluation import (
    REFERENCE_FIELDS as DEVELOPMENT_REFERENCE_FIELDS,
    REFERENCE_TYPE as DEVELOPMENT_REFERENCE_TYPE,
    REVIEW_HEADERS,
)
from app.benchmarks.r18g_group_human_evidence import PARTITION_IDS
from app.benchmarks.r18g_targeted_partition_completion import (
    INTERNAL_MAPPING_FIELDS,
    MEMBER_HEADERS,
    TARGET_IDS,
    prepare_targeted_package,
)
from app.benchmarks.r18g_targeted_partition_evaluation import (
    EXPECTED_ASSIGNMENTS,
    REFERENCE_TYPE,
    TargetedPartitionReferenceError,
    canonicalize_partition,
    co_membership_rows,
    compare_alternatives,
    final_decision,
    partition_metrics,
    purity,
    qualitative_comment_state,
    validate_targeted_submission,
)


def _write_csv(path, fields, rows):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _source_files(tmp_path):
    development = tmp_path / "development.xlsx"
    reference = tmp_path / "development_reference.csv"
    mapping = tmp_path / "development_mapping.csv"
    targeted_mapping = tmp_path / "targeted_mapping.csv"
    workbook = Workbook()
    review = workbook.active
    review.title = "Group Review"
    review.append(REVIEW_HEADERS)
    members = workbook.create_sheet("Members")
    members.append(MEMBER_HEADERS)
    reference_rows = []
    mapping_rows = []
    targeted_rows = []
    comments = {
        TARGET_IDS[0]: "Three Manufactured records versus two inputs that are different objects.",
        TARGET_IDS[1]: "Different revenue streams.",
    }
    for group_index, group_id in enumerate(TARGET_IDS):
        parts = list(EXPECTED_ASSIGNMENTS[group_id])
        raw = "MIXED_GROUP_MULTIPLE_IDENTITIES" if group_index == 0 else "NOT_ONE_IDENTITY"
        reason = "MIXED_GROUP_MULTIPLE_IDENTITIES" if group_index == 0 else "OTHER"
        basis = "MIXED" if group_index == 0 else "UNCLEAR"
        review.append([group_id, len(parts), raw, "MEDIUM", reason, comments[group_id], basis])
        refs = []
        for number, part in enumerate(parts, start=1):
            refs.append(f"ref-{group_index}-{number}")
            members.append([
                group_id, number, part, "", f"Description {number}", "", "", "",
                "pcs", "Part type", "SITE",
            ])
        reference_rows.append({
            "review_group_id": group_id, "member_count": len(parts),
            "raw_group_label": raw, "normalized_group_label": "NOT_ONE_IDENTITY",
            "confidence": "MEDIUM", "reason_code": reason,
            "reviewer_comment": comments[group_id], "comment_basis": basis,
            "partition_detail_status": "UNAVAILABLE",
            "reference_type": DEVELOPMENT_REFERENCE_TYPE,
            "submission_sha256": "a" * 64,
            "dataset_version": "r18g-group-human-evidence-v1",
            "normalization_reason": "",
        })
        source = {
            "review_group_id": group_id,
            "hypothesis_fingerprint": f"fingerprint-{group_index}",
            "source_kind": "ADR_PARTITION_CHALLENGER",
            "current_or_shadow": "SHADOW_ADR_ORDER",
            "current_group_id": "", "work_unit_id": f"work-{group_index}",
            "system_status": "", "group_size": len(parts), "site_count": 1,
            "evidence_provenance": "MIXED_EVIDENCE", "strong_edge_count": 1,
            "native_review_edge_count": 1,
            "r18c_demoted_review_edge_count": 0, "cannot_link_count": 0,
            "generic_only_state": "false", "partition_sensitive": "true",
            "deferred_context": "false", "conflict_context": "false",
            "sampling_stratum": "TEST", "development_or_holdout": "DEVELOPMENT",
            "source_member_refs": "|".join(refs),
        }
        mapping_rows.append(source)
        targeted_rows.append({
            "review_group_id": group_id,
            "hypothesis_fingerprint": source["hypothesis_fingerprint"],
            "comparison_set_id": f"comparison-{group_index}",
            "current_or_shadow": source["current_or_shadow"],
            "work_unit_id": source["work_unit_id"],
            "source_member_refs": source["source_member_refs"],
            "member_count": len(parts),
        })
    workbook.save(development)
    _write_csv(reference, DEVELOPMENT_REFERENCE_FIELDS, reference_rows)
    _write_csv(mapping, tuple(mapping_rows[0]), mapping_rows)
    _write_csv(targeted_mapping, INTERNAL_MAPPING_FIELDS, targeted_rows)
    return development, reference, mapping, targeted_mapping


def _completed_submission(tmp_path):
    development, reference, mapping, targeted_mapping = _source_files(tmp_path)
    output = tmp_path / "output"
    prepare_targeted_package(development, reference, mapping, output)
    submission = output / "r18g_targeted_partition_completion.xlsx"
    workbook = load_workbook(submission)
    comments = {
        TARGET_IDS[0]: {
            "SD-M-CANDLE2": "Under the assumption of prefix M stands for Manufactured",
            "SD-P-CPURCH2": "Under the assumption of prefix P defines Dissembly Component",
        },
        TARGET_IDS[1]: {
            part: "Similar Product with different Revenue streams"
            for part in EXPECTED_ASSIGNMENTS[TARGET_IDS[1]]
        },
    }
    for row in workbook["Partition"].iter_rows(min_row=2):
        group_id, part = row[0].value, row[2].value
        row[3].value = EXPECTED_ASSIGNMENTS[group_id][part]
        row[4].value = comments[group_id].get(part, "")
    workbook.save(submission)
    return submission, development, reference, mapping, targeted_mapping


def test_targeted_reference_validates_two_groups_and_freezes_raw_values(tmp_path):
    paths = _completed_submission(tmp_path)
    result = validate_targeted_submission(*paths)
    assert tuple(result["groups"]) == TARGET_IDS
    assert sum(len(rows) for rows in result["partition_rows"].values()) == 9
    assert result["submission_bytes"] > 0
    assert len(result["submission_sha256"]) == 64
    assert all(
        row["partition_id"] in PARTITION_IDS
        for rows in result["partition_rows"].values() for row in rows
    )


def test_validation_rejects_partial_formula_foreign_and_holdout_inputs(tmp_path):
    paths = list(_completed_submission(tmp_path))
    workbook = load_workbook(paths[0])
    workbook["Partition"]["D2"] = ""
    workbook.save(paths[0])
    with pytest.raises(TargetedPartitionReferenceError, match="partition ID"):
        validate_targeted_submission(*paths)
    workbook = load_workbook(paths[0])
    workbook["Partition"]["D2"] = "=1+1"
    workbook.save(paths[0])
    with pytest.raises(TargetedPartitionReferenceError, match="formula"):
        validate_targeted_submission(*paths)
    holdout = tmp_path / "sealed_holdout.xlsx"
    with pytest.raises(ValueError, match="HOLDOUT_FIREWALL"):
        validate_targeted_submission(holdout, *paths[1:])


def test_partition_canonicalization_is_label_and_group_local():
    first, blocks, unresolved = canonicalize_partition({"a": "P9", "b": "P9", "c": "P4"})
    relabeled, relabeled_blocks, _ = canonicalize_partition({"a": "P1", "b": "P1", "c": "P20"})
    second_group, _, _ = canonicalize_partition({"x": "P9", "y": "P4"})
    assert blocks == relabeled_blocks == (("a", "b"), ("c",))
    assert first == relabeled == {"a": "H1", "b": "H1", "c": "H2"}
    assert second_group == {"x": "H1", "y": "H2"}
    assert unresolved == ()


def test_human_co_membership_and_purity_contracts():
    assignments = {"a": "P1", "b": "P1", "c": "P2", "d": "UNRESOLVED"}
    rows = co_membership_rows("group", assignments)
    counts = {state: sum(row["human_relationship"] == state for row in rows) for state in {
        row["human_relationship"] for row in rows
    }}
    assert counts == {
        "HUMAN_SAME_IDENTITY_WITHIN_SET": 1,
        "HUMAN_DIFFERENT_IDENTITY_WITHIN_SET": 2,
        "HUMAN_PARTITION_UNRESOLVED": 3,
    }
    canonical = canonicalize_partition(assignments)[0]
    assert purity(("a", "b"), canonical) == "HUMAN_PURE_GROUP"
    assert purity(("a", "c"), canonical) == "HUMAN_MIXED_GROUP"
    assert purity(("a", "d"), canonical) == "HUMAN_PARTIALLY_UNRESOLVED"
    assert purity(("a",), canonical) == "HUMAN_SINGLETON"


def test_exact_refinement_coarsening_and_mixed_classifications():
    human = {"a": "H1", "b": "H1", "c": "H2", "d": "H2"}
    assert partition_metrics((("a", "b"), ("c", "d")), human)["classification"] == "EXACT_HUMAN_PARTITION_MATCH"
    assert partition_metrics((("a",), ("b",), ("c", "d")), human)["classification"] == "SAFE_REFINEMENT"
    assert partition_metrics((("a", "b", "c", "d"),), human)["classification"] == "UNSAFE_COARSENING"
    assert partition_metrics((("a", "c"), ("b",), ("d",)), human)["classification"] == "MIXED_REFINEMENT_AND_COARSENING"


def test_comparison_and_final_decision_are_transparent_and_conservative():
    human = {"a": "H1", "b": "H1", "c": "H2"}
    current = partition_metrics((("a", "b"), ("c",)), human)
    adr = partition_metrics((("a", "b", "c"),), human)
    assert compare_alternatives(current, adr) == "CURRENT_STRONGLY_PREFERRED"
    result = [{
        "comparison_decision": "CURRENT_STRONGLY_PREFERRED",
        "current_metrics": current, "adr_metrics": adr,
    }]
    assert final_decision(result) == "R18G_B2_CURRENT_PARTITION_OBJECTIVE_SUPPORTED"
    result.append({
        "comparison_decision": "ADR_SHADOW_STRONGLY_PREFERRED",
        "current_metrics": adr, "adr_metrics": current,
    })
    assert final_decision(result) == "R18G_B2_PARTITION_OBJECTIVE_CHANGE_NOT_JUSTIFIED"


def test_qualitative_tension_never_changes_partition():
    assignments = EXPECTED_ASSIGNMENTS[TARGET_IDS[0]]
    before = dict(assignments)
    state = qualitative_comment_state(
        TARGET_IDS[0], "The two inputs are different objects.", assignments
    )
    assert state == "QUALITATIVE_COMMENT_PARTITION_TENSION"
    assert assignments == before
    assert REFERENCE_TYPE == "SINGLE_SENIOR_DOMAIN_EXPERT_TARGETED_PARTITION_DEVELOPMENT"
