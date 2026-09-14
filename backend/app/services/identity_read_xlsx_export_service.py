"""Client-friendly XLSX representation of authority-selected System Groups."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from app.db.models import DuplicateScan
from app.identity_read.key_codec import serialize_versioned_identity_group_key
from app.services.identity_group_review_service import VersionedIdentityGroupReviewService
from app.services.identity_group_presentation import (
    human_review_presentation,
    system_evidence_tier,
)
from app.services.identity_read_export_service import authority_selected_system_group_rows


WORKBOOK_NOTICE = (
    "System-generated groups are suggestions requiring human review. "
    "They are not automatic merge instructions."
)
SHEET_ORDER = (
    "Overview",
    "Review Groups",
    "Group Index",
    "Detailed Data",
    "Technical Reference",
)
GROUP_INDEX_COLUMNS = (
    "Group", "Review Status", "Evidence", "Members", "Sites",
    "Why Suggested", "Human Decision", "Human Comment",
)
REVIEW_GROUP_COLUMNS = (
    "Group", "Review Status", "Evidence", "Group Sites", "Why Suggested",
    "Human Decision", "Human Comment", "Member #", "Part Number",
    "Description", "Site", "UOM", "Part Type", "Commodity Group 01",
    "Commodity Group 02", "Safety Code", "Accounting Group", "Product Code",
    "Product Family", "Product Category", "HSN/SAC Code",
)
DETAILED_DATA_COLUMNS = (
    "Group", "Review Status", "Evidence", "Members", "Group Sites",
    "Human Decision", "Human Comment", "Part Number", "Description", "Site",
    "Inventory UOM", "Part Type", "Commodity Group 01", "Commodity Group 02",
    "Safety Code", "Accounting Group", "Product Code", "Product Family",
    "Product Category", "HSN/SAC Code",
)
TECHNICAL_REFERENCE_COLUMNS = (
    "Group", "Canonical Group ID", "Member #", "Part Number", "Source Row",
    "Stable Record Reference", "Projection Contract", "Source Projection Run",
    "Original System Reason",
)

# Backward-compatible imports now describe the corresponding client sheets.
GROUP_COLUMNS = GROUP_INDEX_COLUMNS
ALL_COLUMNS = DETAILED_DATA_COLUMNS

_SOURCE_COLUMNS = (
    "Part Number", "Description", "Site", "Inventory UOM", "Part Type",
    "Commodity Group 01", "Commodity Group 02", "Safety Code",
    "Accounting Group", "Product Code", "Product Family", "Product Category",
    "HSN/SAC Code",
)
_MEMBER_FIELD_BY_COLUMN = {
    "Part Number": "part_no",
    "Description": "description",
    "Site": "site_or_contract",
    "Inventory UOM": "uom",
    "Part Type": "part_type",
    "Commodity Group 01": "commodity_group_01",
    "Commodity Group 02": "commodity_group_02",
    "Safety Code": "safety_code",
    "Accounting Group": "accounting_group",
    "Product Code": "product_code",
    "Product Family": "product_family",
    "Product Category": "product_category",
    "HSN/SAC Code": "hsn_sac",
}
_REASONS = {
    "LIKELY_DUPLICATE_GROUP": (
        "Stronger deterministic evidence caused the system to suggest this group "
        "for human review."
    ),
    "POSSIBLE_DUPLICATE_GROUP_REVIEW": (
        "Deterministic review evidence caused the system to suggest this group; "
        "human review is required."
    ),
    "CONFLICT": (
        "Conflicting identity evidence prevents safe grouping; human review is required."
    ),
    "DEFERRED": (
        "Identity evaluation is incomplete or deferred; no same-identity conclusion is implied."
    ),
}
_SHORT_REASONS = {
    "LIKELY_DUPLICATE_GROUP": "Multiple deterministic identity signals support review.",
    "POSSIBLE_DUPLICATE_GROUP_REVIEW": (
        "Some identity signals match; manual assessment is needed."
    ),
    "CONFLICT": "Identity signals conflict; manual assessment is needed.",
    "DEFERRED": "Identity evaluation is incomplete; manual assessment is needed.",
}

_NAVY = "1F4E78"
_PALE_BLUE = "EAF3F8"
_PALE_GRAY = "F3F5F7"
_PALE_GOLD = "FFF2CC"
_PALE_GREEN = "E2F0D9"
_PALE_RED = "FCE4D6"
_PALE_PURPLE = "E4DFEC"
_WHITE = "FFFFFF"
_TEXT = "243746"
_HEADER_FILL = PatternFill("solid", fgColor=_NAVY)
_HEADER_FONT = Font(color=_WHITE, bold=True)
_GROUP_FILLS = (
    PatternFill("solid", fgColor=_PALE_BLUE),
    PatternFill("solid", fgColor=_PALE_GRAY),
)
_THIN_GRAY = Side(style="thin", color="B7C9D6")
_MEDIUM_BLUE = Side(style="medium", color="6C8FA5")
_BORDER = Border(
    left=_THIN_GRAY, right=_THIN_GRAY, top=_THIN_GRAY, bottom=_THIN_GRAY
)


def write_spreadsheet_safe_cell(cell, value) -> None:
    """Write business-controlled text literally, never as an Excel formula."""
    if value is None:
        cell.value = ""
        return
    if isinstance(value, (datetime, date)):
        value = value.isoformat()
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def reason_for_group_status(status: str) -> str:
    return _REASONS.get(
        status, "System-generated same-identity candidate requiring human review."
    )


def concise_reason_for_group_status(status: str) -> str:
    return _SHORT_REASONS.get(
        status, "System suggestion requires manual identity assessment."
    )


def _sites(member_rows) -> str:
    values = sorted({
        str(row.get("site_or_contract") or "").strip()
        for row in member_rows
        if str(row.get("site_or_contract") or "").strip()
    })
    return ", ".join(values) or "Not provided"


def _group_presentation(label: str, group, state: dict | None, member_rows) -> dict:
    review_state, human_decision = human_review_presentation(state)
    return {
        "label": label,
        "canonical_id": serialize_versioned_identity_group_key(group.versioned_group_key),
        "review_status": review_state,
        "evidence": system_evidence_tier(group.status.value),
        "members": group.member_count,
        "sites": _sites(member_rows),
        "why": concise_reason_for_group_status(group.status.value),
        "original_reason": reason_for_group_status(group.status.value),
        "human_decision": human_decision,
        "human_comment": (state or {}).get("comment") or "",
    }


def _source_values(row: dict) -> tuple:
    return tuple(row.get(_MEMBER_FIELD_BY_COLUMN[column]) for column in _SOURCE_COLUMNS)


def _write_row(sheet, row_number: int, values, *, wrap_columns=()) -> None:
    for column_number, value in enumerate(values, start=1):
        cell = sheet.cell(row=row_number, column=column_number)
        write_spreadsheet_safe_cell(cell, value)
        cell.border = _BORDER
        cell.alignment = Alignment(
            vertical="top", wrap_text=column_number in wrap_columns
        )


def _write_header(sheet, columns) -> None:
    _write_row(sheet, 1, columns, wrap_columns=range(1, len(columns) + 1))
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 34
    sheet.sheet_view.showGridLines = False


def _set_widths(sheet, widths) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def _style_state(cell, review_status: str) -> None:
    color = {
        "Human Confirmed Same-Identity Group": _PALE_GREEN,
        "Human Rejected Candidate": _PALE_RED,
        "Review Deferred": _PALE_PURPLE,
    }.get(review_status, _PALE_GOLD)
    cell.fill = PatternFill("solid", fgColor=color)
    cell.font = Font(bold=True, color=_TEXT)


def _merge_and_write(sheet, cell_range: str, value, *, fill, font, alignment) -> None:
    sheet.merge_cells(cell_range)
    cell = sheet[cell_range.split(":", 1)[0]]
    write_spreadsheet_safe_cell(cell, value)
    cell.fill = fill
    cell.font = font
    cell.alignment = alignment
    for row in sheet[cell_range]:
        for item in row:
            item.fill = fill
            item.border = _BORDER


def _write_kpi(sheet, columns: str, label: str, value, *, start_row: int) -> None:
    start_col, end_col = columns.split(":")
    label_range = f"{start_col}{start_row}:{end_col}{start_row}"
    value_range = f"{start_col}{start_row + 1}:{end_col}{start_row + 2}"
    _merge_and_write(
        sheet, label_range, label,
        fill=PatternFill("solid", fgColor=_NAVY),
        font=Font(color=_WHITE, bold=True),
        alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
    )
    _merge_and_write(
        sheet, value_range, value,
        fill=PatternFill("solid", fgColor=_PALE_BLUE if start_row == 7 else _PALE_GRAY),
        font=Font(color=_NAVY, bold=True, size=18),
        alignment=Alignment(horizontal="center", vertical="center"),
    )


def _write_overview(workbook, scan, snapshot, review_states) -> None:
    sheet = workbook.active
    sheet.title = "Overview"
    sheet.sheet_view.showGridLines = False
    _set_widths(sheet, (18,) * 8)
    _merge_and_write(
        sheet, "A1:H2", "Inventory Identity Review Candidate Report",
        fill=PatternFill("solid", fgColor=_NAVY),
        font=Font(color=_WHITE, bold=True, size=20),
        alignment=Alignment(horizontal="center", vertical="center"),
    )
    _merge_and_write(
        sheet, "A4:H5",
        WORKBOOK_NOTICE + " Human decisions are authoritative and override system suggestions.",
        fill=PatternFill("solid", fgColor=_PALE_GOLD),
        font=Font(color=_TEXT, bold=True),
        alignment=Alignment(horizontal="left", vertical="center", wrap_text=True),
    )
    confirmed = sum(
        state.get("current_decision_type")
        in {"CONFIRM_ALL_AS_ONE", "CONFIRM_SELECTED", "SPLIT_PARTITIONS"}
        for state in review_states.values()
    )
    rejected = sum(
        state.get("current_decision_type") == "KEEP_ALL_SEPARATE"
        for state in review_states.values()
    )
    deferred = sum(
        not state.get("reviewed") or state.get("current_decision_type") == "UNSURE"
        for state in review_states.values()
    )
    for columns, label, value in (
        ("A:B", "Input Records", snapshot.canonical_record_count),
        ("C:D", "Candidate Groups", snapshot.group_count),
        ("E:F", "Stronger Evidence", snapshot.likely_group_count),
        ("G:H", "Needs Additional Review", snapshot.review_group_count),
    ):
        _write_kpi(sheet, columns, label, value, start_row=7)
    for columns, label, value in (
        ("A:B", "Human Confirmed", confirmed),
        ("C:D", "Human Rejected", rejected),
        ("E:F", "Deferred / Unreviewed", deferred),
        ("G:H", "Conflicts / Deferred Work", snapshot.conflict_count + snapshot.deferred_count),
    ):
        _write_kpi(sheet, columns, label, value, start_row=11)
    _merge_and_write(
        sheet, "A15:H15", "How to use this workbook",
        fill=PatternFill("solid", fgColor=_NAVY),
        font=Font(color=_WHITE, bold=True),
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    _merge_and_write(
        sheet, "A16:H19",
        "1. Start with Review Groups.\n"
        "2. Review the records inside each candidate group.\n"
        "3. System evidence is advisory.\n"
        "4. Human decisions are authoritative.",
        fill=PatternFill("solid", fgColor=_WHITE),
        font=Font(color=_TEXT),
        alignment=Alignment(horizontal="left", vertical="top", wrap_text=True),
    )
    metadata = (
        ("Scan ID", snapshot.scan_id),
        ("Scan Name", scan.scan_name),
        ("Scan Status", scan.status),
        ("Projection Contract", snapshot.projection_contract.value),
        ("Source Projection Run", snapshot.source_projection_run_id),
        ("Unassigned Records", snapshot.unassigned_count),
    )
    _merge_and_write(
        sheet, "A21:H21", "Report details",
        fill=PatternFill("solid", fgColor="5B7894"),
        font=Font(color=_WHITE, bold=True),
        alignment=Alignment(horizontal="left", vertical="center"),
    )
    for row_number, (label, value) in enumerate(metadata, start=22):
        _merge_and_write(
            sheet, f"A{row_number}:B{row_number}", label,
            fill=PatternFill("solid", fgColor=_PALE_GRAY),
            font=Font(color=_TEXT, bold=True),
            alignment=Alignment(vertical="center"),
        )
        _merge_and_write(
            sheet, f"C{row_number}:H{row_number}", value,
            fill=PatternFill("solid", fgColor=_WHITE),
            font=Font(color=_TEXT),
            alignment=Alignment(vertical="center", wrap_text=True),
        )
    sheet.freeze_panes = "A7"
    for row_number in (1, 2, 4, 5, 8, 9, 12, 13):
        sheet.row_dimensions[row_number].height = 26


def _write_group_index(sheet, groups) -> None:
    _write_header(sheet, GROUP_INDEX_COLUMNS)
    for row_number, item in enumerate(groups, start=2):
        p = item["presentation"]
        _write_row(
            sheet, row_number,
            (p["label"], p["review_status"], p["evidence"], p["members"],
             p["sites"], p["why"], p["human_decision"], p["human_comment"]),
            wrap_columns=(2, 5, 6, 7, 8),
        )
        _style_state(sheet.cell(row_number, 2), p["review_status"])
        sheet.row_dimensions[row_number].height = 36
    _set_widths(sheet, (16, 34, 20, 11, 22, 45, 32, 45))
    if groups:
        sheet.auto_filter.ref = f"A1:H{sheet.max_row}"


def _write_review_groups(sheet, groups) -> None:
    _write_header(sheet, REVIEW_GROUP_COLUMNS)
    current_row = 2
    if not groups:
        _merge_and_write(
            sheet, "A2:U3", "No candidate groups were generated for this scan.",
            fill=PatternFill("solid", fgColor=_PALE_GRAY),
            font=Font(color=_TEXT, italic=True),
            alignment=Alignment(horizontal="center", vertical="center"),
        )
    for group_index, item in enumerate(groups):
        p = item["presentation"]
        member_rows = item["member_rows"]
        start_row = current_row
        for member_number, member_row in enumerate(member_rows, start=1):
            source = _source_values(member_row)
            _write_row(
                sheet, current_row,
                (p["label"], p["review_status"], p["evidence"], p["sites"],
                 p["why"], p["human_decision"], p["human_comment"],
                 member_number, *source),
                wrap_columns=(2, 4, 5, 6, 7, 10),
            )
            sheet.row_dimensions[current_row].height = 42
            current_row += 1
        end_row = current_row - 1
        fill = _GROUP_FILLS[group_index % len(_GROUP_FILLS)]
        for row_number in range(start_row, end_row + 1):
            for column_number in range(1, len(REVIEW_GROUP_COLUMNS) + 1):
                cell = sheet.cell(row_number, column_number)
                cell.fill = fill
                cell.border = Border(
                    left=_THIN_GRAY, right=_THIN_GRAY,
                    top=_MEDIUM_BLUE if row_number == start_row else _THIN_GRAY,
                    bottom=_MEDIUM_BLUE if row_number == end_row else _THIN_GRAY,
                )
        for column_number in range(1, 8):
            if end_row > start_row:
                sheet.merge_cells(
                    start_row=start_row, start_column=column_number,
                    end_row=end_row, end_column=column_number,
                )
            sheet.cell(start_row, column_number).alignment = Alignment(
                vertical="center", wrap_text=True
            )
        _style_state(sheet.cell(start_row, 2), p["review_status"])
    _set_widths(
        sheet,
        (16, 34, 20, 22, 42, 32, 42, 10, 20, 48, 18, 14, 18, 22, 22,
         16, 20, 18, 20, 20, 18),
    )


def _write_detailed_data(sheet, groups) -> None:
    _write_header(sheet, DETAILED_DATA_COLUMNS)
    row_number = 2
    for item in groups:
        p = item["presentation"]
        for member_row in item["member_rows"]:
            _write_row(
                sheet, row_number,
                (p["label"], p["review_status"], p["evidence"], p["members"],
                 p["sites"], p["human_decision"], p["human_comment"],
                 *_source_values(member_row)),
                wrap_columns=(2, 5, 6, 7, 9),
            )
            _style_state(sheet.cell(row_number, 2), p["review_status"])
            sheet.row_dimensions[row_number].height = 36
            row_number += 1
    _set_widths(
        sheet,
        (16, 34, 20, 11, 22, 32, 42, 20, 48, 18, 16, 18, 22, 22, 16,
         20, 18, 20, 20, 18),
    )
    if sheet.max_row >= 2:
        table = Table(
            displayName="SystemGroupData",
            ref=f"A1:{get_column_letter(len(DETAILED_DATA_COLUMNS))}{sheet.max_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)


def _write_technical_reference(sheet, groups, snapshot) -> None:
    _write_header(sheet, TECHNICAL_REFERENCE_COLUMNS)
    row_number = 2
    for item in groups:
        p = item["presentation"]
        for member_number, member_row in enumerate(item["member_rows"], start=1):
            _write_row(
                sheet, row_number,
                (p["label"], p["canonical_id"], member_number,
                 member_row.get("part_no"), member_row.get("source_row_reference"),
                 member_row.get("stable_record_reference"),
                 snapshot.projection_contract.value,
                 snapshot.source_projection_run_id, p["original_reason"]),
                wrap_columns=(2, 6, 9),
            )
            row_number += 1
    _set_widths(sheet, (16, 38, 10, 20, 14, 42, 22, 22, 58))
    if groups:
        sheet.auto_filter.ref = f"A1:I{sheet.max_row}"


def authority_selected_system_groups_to_xlsx(db, scan_id: int) -> bytes:
    """Create a client workbook from the unchanged System Group projection."""
    snapshot, export_rows = authority_selected_system_group_rows(db, scan_id)
    scan = db.get(DuplicateScan, scan_id)
    review_states = VersionedIdentityGroupReviewService(db).current_states_for_snapshot(
        snapshot
    )
    rows_by_group = {}
    for row in export_rows:
        rows_by_group.setdefault(row["group_key"], []).append(row)

    groups = []
    for group_index, group in enumerate(snapshot.groups, start=1):
        canonical_id = serialize_versioned_identity_group_key(group.versioned_group_key)
        member_rows = tuple(rows_by_group.get(canonical_id, ()))
        groups.append({
            "presentation": _group_presentation(
                f"CG-{group_index:06d}", group,
                review_states.get(canonical_id), member_rows,
            ),
            "member_rows": member_rows,
        })

    workbook = Workbook()
    _write_overview(workbook, scan, snapshot, review_states)
    review_groups = workbook.create_sheet("Review Groups")
    group_index = workbook.create_sheet("Group Index")
    detailed_data = workbook.create_sheet("Detailed Data")
    technical = workbook.create_sheet("Technical Reference")
    _write_review_groups(review_groups, groups)
    _write_group_index(group_index, groups)
    _write_detailed_data(detailed_data, groups)
    _write_technical_reference(technical, groups, snapshot)
    workbook.active = 0

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _write_reviewed_summary(workbook, scan, snapshot, rows) -> None:
    sheet = workbook.active
    sheet.title = "Summary"
    set_count = len({row["reviewed_identity_set_key"] for row in rows})
    summary_rows = (
        ("Reviewed Identity Export", ""),
        ("Notice", REVIEWED_WORKBOOK_NOTICE),
        ("Report type", "Reviewed Identity Export"),
        ("Authority", "Human-confirmed"),
        ("Human confirmation", "Required and applied"),
        ("Scan identifier", snapshot.scan_id),
        ("Scan name", scan.scan_name),
        ("Scan status", scan.status),
        ("Input record count", snapshot.canonical_record_count),
        ("Reviewed identity set count", set_count),
        ("Reviewed member count", len(rows)),
        ("Projection contract", snapshot.projection_contract.value),
        ("Source projection run", snapshot.source_projection_run_id),
    )
    for row_number, values in enumerate(summary_rows, start=1):
        _write_row(sheet, row_number, values)
    sheet["A1"].font = Font(bold=True, size=16, color="1F4E78")
    for row_number in range(2, len(summary_rows) + 1):
        sheet.cell(row_number, 1).font = Font(bold=True)
    sheet["B2"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 100
    sheet.freeze_panes = "A2"


def _reviewed_set_values(index: int, set_rows: list[dict]) -> tuple:
    first_row = set_rows[0]
    decision = first_row["review_decision_type"]
    return (
        f"RS-{index:06d}",
        first_row["group_reference"],
        _REVIEW_LABELS.get(decision, decision),
        first_row["reviewer"],
        first_row["reviewed_at"],
        first_row.get("review_comment") or "",
        len(set_rows),
    )


def authority_selected_reviewed_identities_to_xlsx(db, scan_id: int) -> bytes:
    """Create an in-memory workbook from the exact Reviewed Identity export projection."""
    snapshot, rows = authority_selected_reviewed_identity_rows(db, scan_id)
    scan = db.get(DuplicateScan, scan_id)

    rows_by_set: dict[str, list[dict]] = {}
    set_order: list[str] = []
    for row in rows:
        set_key = row["reviewed_identity_set_key"]
        if set_key not in rows_by_set:
            rows_by_set[set_key] = []
            set_order.append(set_key)
        rows_by_set[set_key].append(row)

    workbook = Workbook()
    _write_reviewed_summary(workbook, scan, snapshot, rows)
    sheet = workbook.create_sheet("Reviewed Identity Sets")
    _write_header(sheet, ALL_REVIEWED_COLUMNS)

    row_number = 2
    for set_index, set_key in enumerate(set_order, start=1):
        set_rows = rows_by_set[set_key]
        set_values = _reviewed_set_values(set_index, set_rows)
        first_data_row = row_number
        for member_row in set_rows:
            _write_row(sheet, row_number, set_values + _member_values(member_row))
            row_number += 1
        last_data_row = row_number - 1
        if last_data_row > first_data_row:
            for column in range(1, len(REVIEWED_SET_COLUMNS) + 1):
                sheet.merge_cells(
                    start_row=first_data_row,
                    start_column=column,
                    end_row=last_data_row,
                    end_column=column,
                )
                sheet.cell(first_data_row, column).alignment = Alignment(
                    vertical="top", wrap_text=column in (3, 6)
                )
        for row_index in range(first_data_row, last_data_row + 1):
            for column in range(1, len(REVIEWED_SET_COLUMNS) + 1):
                sheet.cell(row_index, column).fill = _GROUP_FILL

    _style_dimensions(sheet, _REVIEWED_WIDTHS)
    last_column = get_column_letter(len(ALL_REVIEWED_COLUMNS))
    sheet.auto_filter.ref = f"A1:{last_column}{max(1, sheet.max_row)}"
    if sheet.max_row >= 2:
        table = Table(
            displayName="ReviewedIdentitySets", ref=f"A1:{last_column}{sheet.max_row}"
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
            showRowStripes=True, showColumnStripes=False,
        )
        sheet.add_table(table)

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
