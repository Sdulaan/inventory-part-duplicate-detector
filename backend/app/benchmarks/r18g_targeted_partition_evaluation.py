"""Evaluate the completed R18G-B1 targeted Development partitions offline."""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook

from app.benchmarks.r18g_development_group_evaluation import assert_holdout_firewall
from app.benchmarks.r18g_group_human_evidence import (
    PARTITION_IDS,
    _pair,
    _work_unit_members,
    fingerprint,
    load_current_resolution,
    sha256_file,
    write_manifest,
)
from app.benchmarks.r18g_targeted_partition_completion import (
    EXPECTED_MEMBER_COUNTS,
    FORBIDDEN_HEADERS,
    GROUP_HEADERS,
    INTERNAL_MAPPING_FIELDS,
    MEMBER_HEADERS,
    PARTITION_HEADERS,
    TARGET_IDS,
    load_target_sources,
)


REFERENCE_TYPE = "SINGLE_SENIOR_DOMAIN_EXPERT_TARGETED_PARTITION_DEVELOPMENT"
HUMAN_EVIDENCE_STRENGTH = "MEDIUM_SINGLE_SENIOR_EXPERT"
EXPECTED_ASSIGNMENTS = {
    "R18G-D09438A953125080": {
        "SD-M-CANDLE2": "P1",
        "SD-CANDLE": "P1",
        "SD-P-CPURCH2": "P2",
        "SD-M-CANDLE": "P1",
        "SD-P-WAX3": "P2",
    },
    "R18G-FFF38D9F3046F48C": {
        "MLR-SAL2-02.25.2023": "P3",
        "MLR-PUR-02.25.2023": "P4",
        "MLR-MANU-02.25.20232": "P5",
        "MLR-INV-02.25.2023": "P6",
    },
}
EXPECTED_COMMENT_FRAGMENTS = {
    "R18G-D09438A953125080": {
        "Under the assumption of prefix M stands for Manufactured",
        "Under the assumption of prefix P defines Dissembly Component",
    },
    "R18G-FFF38D9F3046F48C": {
        "Similar Product with different Revenue streams",
    },
}
REFERENCE_FIELDS = (
    "review_group_id", "member_no", "part_number", "human_partition_id_raw",
    "canonical_human_partition", "partition_comment",
    "previous_raw_group_label", "previous_normalized_group_label",
    "previous_confidence", "reference_type", "submission_sha256",
)
CANDIDATE_FIELDS = (
    "review_group_id", "comparison_set_id", "alternative", "candidate_group_no",
    "candidate_member_refs", "candidate_part_numbers", "member_count", "purity",
    "cannot_link_count",
)


class TargetedPartitionReferenceError(ValueError):
    """Fail closed when the targeted human reference or comparison is invalid."""


def _csv_rows(path: Path):
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _rows(sheet, width):
    return [
        tuple(row[index].value for index in range(width))
        for row in sheet.iter_rows(min_row=2)
        if any(row[index].value not in (None, "") for index in range(width))
    ]


def canonicalize_partition(assignments):
    """Return label-independent canonical blocks ordered by their member keys."""
    unresolved = tuple(sorted(member for member, label in assignments.items()
                              if label == "UNRESOLVED"))
    grouped = {}
    for member, label in assignments.items():
        if label == "UNRESOLVED":
            continue
        grouped.setdefault(label, set()).add(member)
    blocks = tuple(sorted((tuple(sorted(block)) for block in grouped.values())))
    canonical = {
        member: f"H{index}" for index, block in enumerate(blocks, start=1)
        for member in block
    }
    canonical.update({member: "UNRESOLVED" for member in unresolved})
    return canonical, blocks, unresolved


def co_membership_rows(group_id, assignments):
    canonical, _, unresolved = canonicalize_partition(assignments)
    unresolved = set(unresolved)
    rows = []
    for left, right in combinations(sorted(assignments), 2):
        if left in unresolved or right in unresolved:
            state = "HUMAN_PARTITION_UNRESOLVED"
        elif canonical[left] == canonical[right]:
            state = "HUMAN_SAME_IDENTITY_WITHIN_SET"
        else:
            state = "HUMAN_DIFFERENT_IDENTITY_WITHIN_SET"
        rows.append({"review_group_id": group_id, "left_member_ref": left,
                     "right_member_ref": right, "human_relationship": state})
    return rows


def validate_targeted_submission(submission_path: Path, development_path: Path,
                                 reference_path: Path, source_mapping_path: Path,
                                 targeted_mapping_path: Path):
    assert_holdout_firewall(
        submission_path, development_path, reference_path, source_mapping_path,
        targeted_mapping_path,
    )
    expected_groups, _, expected_members = load_target_sources(
        development_path, reference_path, source_mapping_path
    )
    targeted_mapping = _csv_rows(targeted_mapping_path)
    if (len(targeted_mapping) != 2
            or tuple(targeted_mapping[0]) != INTERNAL_MAPPING_FIELDS
            or tuple(row["review_group_id"] for row in targeted_mapping) != TARGET_IDS):
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: mapping")

    workbook = load_workbook(submission_path, data_only=False, read_only=True)
    if workbook.sheetnames != ["Instructions", "Groups", "Members", "Partition"]:
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: sheets")
    if tuple(cell.value for cell in workbook["Groups"][1]) != GROUP_HEADERS:
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: groups schema")
    if tuple(cell.value for cell in workbook["Members"][1]) != MEMBER_HEADERS:
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: members schema")
    if tuple(cell.value for cell in workbook["Partition"][1]) != PARTITION_HEADERS:
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: partition schema")

    expected_group_rows = tuple(
        (
            group_id, int(expected_groups[group_id]["member_count"]),
            expected_groups[group_id]["raw_group_label"],
            expected_groups[group_id]["normalized_group_label"],
            expected_groups[group_id]["confidence"],
            expected_groups[group_id]["reason_code"],
            expected_groups[group_id]["reviewer_comment"],
            expected_groups[group_id]["comment_basis"],
        ) for group_id in TARGET_IDS
    )
    if tuple(_rows(workbook["Groups"], len(GROUP_HEADERS))) != expected_group_rows:
        raise TargetedPartitionReferenceError(
            "R18G_B2_TARGETED_REFERENCE_INVALID: prior judgment changed"
        )
    if tuple(_rows(workbook["Members"], len(MEMBER_HEADERS))) != expected_members:
        raise TargetedPartitionReferenceError(
            "R18G_B2_TARGETED_REFERENCE_INVALID: source/member fields changed"
        )

    partition_cells = list(workbook["Partition"].iter_rows(min_row=2))
    if len(partition_cells) != 9:
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: row count")
    expected_keys = {(row[0], int(row[1]), row[2]) for row in expected_members}
    actual_keys = {
        (row[0].value, int(row[1].value), row[2].value) for row in partition_cells
    }
    if actual_keys != expected_keys:
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: members")
    if any(row[3].data_type == "f" for row in partition_cells):
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: formula")
    if any(str(row[3].value or "").strip() not in PARTITION_IDS for row in partition_cells):
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: partition ID")
    visible_headers = {
        str(cell.value or "").casefold() for sheet in workbook for cell in sheet[1]
    }
    if visible_headers & FORBIDDEN_HEADERS:
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: leakage")
    if any("holdout" in str(cell.value or "").casefold()
           for sheet in workbook for row in sheet.iter_rows() for cell in row):
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: holdout")

    by_group = {group_id: [] for group_id in TARGET_IDS}
    for row in partition_cells:
        by_group[row[0].value].append({
            "member_no": int(row[1].value), "part_number": str(row[2].value),
            "partition_id": str(row[3].value).strip(),
            "comment": str(row[4].value or "").strip(),
        })
    for group_id, values in by_group.items():
        if len(values) != EXPECTED_MEMBER_COUNTS[group_id]:
            raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: count")
        observed = {row["part_number"]: row["partition_id"] for row in values}
        if observed != EXPECTED_ASSIGNMENTS[group_id]:
            raise TargetedPartitionReferenceError(
                "R18G_B2_TARGETED_REFERENCE_INVALID: transcription mismatch"
            )
        comments = {row["comment"] for row in values if row["comment"]}
        if not EXPECTED_COMMENT_FRAGMENTS[group_id] <= comments:
            raise TargetedPartitionReferenceError(
                "R18G_B2_TARGETED_REFERENCE_INVALID: comment mismatch"
            )

    return {
        "submission_sha256": sha256_file(submission_path),
        "submission_bytes": submission_path.stat().st_size,
        "groups": expected_groups,
        "members": expected_members,
        "partition_rows": by_group,
        "targeted_mapping": {row["review_group_id"]: row for row in targeted_mapping},
    }


def ranked_partition_alternatives(unit, edges):
    """Reproduce the frozen current/ADR objective comparison transparently."""
    unit = tuple(sorted(unit))
    candidates = {}
    for size in range(2, len(unit) + 1):
        for subset in combinations(unit, size):
            values = [edges.get(_pair(left, right)) for left, right in combinations(subset, 2)]
            if not values or not all(
                edge is not None and edge.edge_class in {"STRONG_SUPPORT", "REVIEW_SUPPORT"}
                for edge in values
            ):
                continue
            if size >= 3 and all(edge.generic_only for edge in values):
                continue
            strong = sum(edge.edge_class == "STRONG_SUPPORT" for edge in values)
            review = sum(edge.edge_class == "REVIEW_SUPPORT" for edge in values)
            likely_members = size if review == 0 else 0
            candidates[frozenset(subset)] = (strong, review, likely_members)
    if not candidates:
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: no candidates")
    partitions = []

    def visit(index, blocks):
        if index == len(unit):
            selected = []
            for block in blocks:
                if len(block) == 1:
                    continue
                metrics = candidates.get(frozenset(block))
                if metrics is None:
                    return
                selected.append((frozenset(block), *metrics))
            used = {member for candidate in selected for member in candidate[0]}
            signature = tuple(sorted(tuple(sorted(candidate[0])) for candidate in selected))
            partitions.append((
                len(used), sum(candidate[3] for candidate in selected),
                sum(candidate[1] for candidate in selected),
                sum(candidate[2] for candidate in selected), len(selected), signature,
            ))
            return
        member = unit[index]
        for block_index in range(len(blocks)):
            updated = [list(block) for block in blocks]
            updated[block_index].append(member)
            visit(index + 1, updated)
        visit(index + 1, [*blocks, [member]])

    visit(0, [])
    current = max(partitions, key=lambda item: (
        item[0], item[1], item[2], -item[3], -item[4], item[5]
    ))
    adr = max(partitions, key=lambda item: (
        item[0], item[1], item[2], item[3], -item[4], item[5]
    ))

    def complete(signature):
        used = {member for block in signature for member in block}
        return tuple(sorted((*signature, *((member,) for member in unit if member not in used))))

    return complete(current[5]), complete(adr[5])


def purity(block, human_by_member):
    if len(block) == 1:
        return "HUMAN_SINGLETON"
    labels = {human_by_member[member] for member in block}
    if "UNRESOLVED" in labels:
        return "HUMAN_PARTIALLY_UNRESOLVED"
    return "HUMAN_PURE_GROUP" if len(labels) == 1 else "HUMAN_MIXED_GROUP"


def partition_metrics(system_blocks, human_by_member):
    members = set(human_by_member)
    if {member for block in system_blocks for member in block} != members:
        raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: coverage")
    same_pairs = {
        frozenset((left, right)) for left, right in combinations(sorted(members), 2)
        if human_by_member[left] != "UNRESOLVED"
        and human_by_member[left] == human_by_member[right]
    }
    different_pairs = {
        frozenset((left, right)) for left, right in combinations(sorted(members), 2)
        if "UNRESOLVED" not in {human_by_member[left], human_by_member[right]}
        and human_by_member[left] != human_by_member[right]
    }
    system_pairs = {
        frozenset((left, right)) for block in system_blocks
        for left, right in combinations(block, 2)
    }
    mixed = sum(purity(block, human_by_member) == "HUMAN_MIXED_GROUP"
                for block in system_blocks)
    pure = sum(len(block) >= 2 and purity(block, human_by_member) == "HUMAN_PURE_GROUP"
               for block in system_blocks)
    preserved = len(same_pairs & system_pairs)
    incorrect = len(different_pairs & system_pairs)
    fragmented = sum(
        len({index for index, block in enumerate(system_blocks)
             if set(block) & human_members}) > 1
        for label in sorted(set(human_by_member.values()) - {"UNRESOLVED"})
        if len(human_members := {m for m, value in human_by_member.items() if value == label}) > 1
    )
    unresolved = sum(value == "UNRESOLVED" for value in human_by_member.values())
    if unresolved:
        relation = "NOT_COMPARABLE"
    else:
        refinement = preserved < len(same_pairs)
        coarsening = incorrect > 0
        if not refinement and not coarsening:
            relation = "EXACT_HUMAN_PARTITION_MATCH"
        elif refinement and not coarsening:
            relation = "SAFE_REFINEMENT"
        elif coarsening and not refinement:
            relation = "UNSAFE_COARSENING"
        else:
            relation = "MIXED_REFINEMENT_AND_COARSENING"
    return {
        "human_mixed_candidate_groups": mixed,
        "human_pure_candidate_groups_size_gte_2": pure,
        "human_same_identity_pairs_preserved_together": preserved,
        "human_same_identity_pairs_total": len(same_pairs),
        "human_different_identity_pairs_incorrectly_co_grouped": incorrect,
        "human_partitions_fragmented": fragmented,
        "unresolved_members": unresolved,
        "coverage": len(members),
        "classification": relation,
    }


def compare_alternatives(current, adr):
    if current["unresolved_members"] or adr["unresolved_members"]:
        return "HUMAN_EVIDENCE_INSUFFICIENT"
    current_key = (
        -current["human_different_identity_pairs_incorrectly_co_grouped"],
        -current["human_mixed_candidate_groups"],
        current["human_same_identity_pairs_preserved_together"],
        -current["human_partitions_fragmented"],
    )
    adr_key = (
        -adr["human_different_identity_pairs_incorrectly_co_grouped"],
        -adr["human_mixed_candidate_groups"],
        adr["human_same_identity_pairs_preserved_together"],
        -adr["human_partitions_fragmented"],
    )
    if current_key == adr_key:
        return "EQUIVALENT_FOR_HUMAN_PARTITION"
    current_unsafe = current["human_different_identity_pairs_incorrectly_co_grouped"] > 0
    adr_unsafe = adr["human_different_identity_pairs_incorrectly_co_grouped"] > 0
    if current_unsafe and not adr_unsafe:
        return ("ADR_SAFER_BUT_MORE_CONSERVATIVE"
                if adr["human_same_identity_pairs_preserved_together"]
                < current["human_same_identity_pairs_preserved_together"]
                else "ADR_SHADOW_STRONGLY_PREFERRED")
    if adr_unsafe and not current_unsafe:
        return ("CURRENT_SAFER_BUT_MORE_CONSERVATIVE"
                if current["human_same_identity_pairs_preserved_together"]
                < adr["human_same_identity_pairs_preserved_together"]
                else "CURRENT_STRONGLY_PREFERRED")
    if current_key < adr_key:
        return "ADR_SHADOW_STRONGLY_PREFERRED"
    if current_key > adr_key:
        return "CURRENT_STRONGLY_PREFERRED"
    return "MIXED_TRADEOFF"


def qualitative_comment_state(group_id, prior_comment, assignments_by_part):
    """Record explicit prose/partition tension without overriding the partition."""
    disassembly_parts = {"SD-P-CPURCH2", "SD-P-WAX3"}
    if (
        group_id == "R18G-D09438A953125080"
        and "different objects" in str(prior_comment).casefold()
        and disassembly_parts <= set(assignments_by_part)
        and len({assignments_by_part[part] for part in disassembly_parts}) == 1
    ):
        return "QUALITATIVE_COMMENT_PARTITION_TENSION"
    return "NONE_OBSERVED"


def final_decision(results):
    outcomes = {item["comparison_decision"] for item in results}
    adr_positive = {"ADR_SHADOW_STRONGLY_PREFERRED", "ADR_SAFER_BUT_MORE_CONSERVATIVE"}
    current_positive = {"CURRENT_STRONGLY_PREFERRED", "CURRENT_SAFER_BUT_MORE_CONSERVATIVE"}
    equivalent = "EQUIVALENT_FOR_HUMAN_PARTITION"
    if outcomes <= adr_positive | {equivalent} and outcomes & adr_positive and all(
        item["adr_metrics"]["human_mixed_candidate_groups"]
        <= item["current_metrics"]["human_mixed_candidate_groups"]
        for item in results
    ):
        return "R18G_B2_ADR_PARTITION_OBJECTIVE_SUPPORTED"
    if outcomes <= current_positive | {equivalent} and outcomes & current_positive:
        return "R18G_B2_CURRENT_PARTITION_OBJECTIVE_SUPPORTED"
    return "R18G_B2_PARTITION_OBJECTIVE_CHANGE_NOT_JUSTIFIED"


def evaluate(submission_path: Path, development_path: Path, reference_path: Path,
             source_mapping_path: Path, targeted_mapping_path: Path,
             database_path: Path, output_dir: Path):
    assert_holdout_firewall(
        submission_path, development_path, reference_path, source_mapping_path,
        targeted_mapping_path, database_path, output_dir,
    )
    validated = validate_targeted_submission(
        submission_path, development_path, reference_path, source_mapping_path,
        targeted_mapping_path,
    )
    value, _, members, edges, source = load_current_resolution(database_path)
    member_by_ref = {member.stable_ref: member for member in members.values()}
    units = _work_unit_members(value)
    unit_by_id = {
        fingerprint((members[item].stable_ref for item in unit), "WORK_UNIT"): unit
        for unit in units
    }
    reference_rows = []
    candidate_rows = []
    comparison_results = []
    co_membership = []

    for group_id in TARGET_IDS:
        mapping = validated["targeted_mapping"][group_id]
        source_refs = tuple(mapping["source_member_refs"].split("|"))
        human_rows = sorted(validated["partition_rows"][group_id], key=lambda row: row["member_no"])
        if len(source_refs) != len(human_rows):
            raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: ref count")
        raw_by_ref = {ref: row["partition_id"] for ref, row in zip(source_refs, human_rows)}
        canonical_by_ref, human_blocks, unresolved = canonicalize_partition(raw_by_ref)
        co_membership.extend(co_membership_rows(group_id, raw_by_ref))
        group_reference = validated["groups"][group_id]
        for ref, row in zip(source_refs, human_rows):
            reference_rows.append({
                "review_group_id": group_id, "member_no": row["member_no"],
                "part_number": row["part_number"],
                "human_partition_id_raw": row["partition_id"],
                "canonical_human_partition": canonical_by_ref[ref],
                "partition_comment": row["comment"],
                "previous_raw_group_label": group_reference["raw_group_label"],
                "previous_normalized_group_label": group_reference["normalized_group_label"],
                "previous_confidence": group_reference["confidence"],
                "reference_type": REFERENCE_TYPE,
                "submission_sha256": validated["submission_sha256"],
            })

        unit = unit_by_id.get(mapping["work_unit_id"])
        if unit is None:
            raise TargetedPartitionReferenceError("R18G_B2_TARGETED_REFERENCE_INVALID: work unit")
        unit_refs = {members[item].stable_ref for item in unit}
        if unit_refs != set(source_refs):
            raise TargetedPartitionReferenceError(
                "R18G_B2_TARGETED_REFERENCE_INVALID: incomplete comparison coverage"
            )
        current_ids, adr_ids = ranked_partition_alternatives(unit, edges)
        target_ids = tuple(sorted(member_by_ref[ref].record_id for ref in source_refs))
        if target_ids not in adr_ids or target_ids in current_ids:
            raise TargetedPartitionReferenceError(
                "R18G_B2_TARGETED_REFERENCE_INVALID: challenger reconciliation"
            )

        def ref_blocks(id_blocks):
            return tuple(sorted(tuple(sorted(members[item].stable_ref for item in block))
                                for block in id_blocks))

        alternatives = {"CURRENT_GF5": ref_blocks(current_ids), "ADR_SHADOW": ref_blocks(adr_ids)}
        metrics = {}
        for alternative, blocks in alternatives.items():
            metrics[alternative] = partition_metrics(blocks, canonical_by_ref)
            for number, block in enumerate(blocks, start=1):
                cannot = sum(
                    edges.get(_pair(member_by_ref[left].record_id, member_by_ref[right].record_id)).edge_class
                    == "CANNOT_LINK"
                    for left, right in combinations(block, 2)
                    if edges.get(_pair(member_by_ref[left].record_id, member_by_ref[right].record_id))
                    is not None
                )
                candidate_rows.append({
                    "review_group_id": group_id,
                    "comparison_set_id": mapping["comparison_set_id"],
                    "alternative": alternative, "candidate_group_no": number,
                    "candidate_member_refs": "|".join(block),
                    "candidate_part_numbers": "|".join(member_by_ref[ref].part_number for ref in block),
                    "member_count": len(block), "purity": purity(block, canonical_by_ref),
                    "cannot_link_count": cannot,
                })
        comparison_decision = compare_alternatives(
            metrics["CURRENT_GF5"], metrics["ADR_SHADOW"]
        )
        tension = qualitative_comment_state(
            group_id, group_reference["reviewer_comment"],
            {row["part_number"]: row["partition_id"] for row in human_rows},
        )
        comparison_results.append({
            "review_group_id": group_id,
            "comparison_set_id": mapping["comparison_set_id"],
            "human_partition_blocks": [list(block) for block in human_blocks],
            "human_unresolved_members": list(unresolved),
            "current_partition_blocks": [list(block) for block in alternatives["CURRENT_GF5"]],
            "adr_shadow_partition_blocks": [list(block) for block in alternatives["ADR_SHADOW"]],
            "current_metrics": metrics["CURRENT_GF5"],
            "adr_metrics": metrics["ADR_SHADOW"],
            "comparison_decision": comparison_decision,
            "qualitative_comment_state": tension,
            "cannot_link_safety": "PRESERVED" if not any(
                row["cannot_link_count"] for row in candidate_rows
                if row["review_group_id"] == group_id
            ) else "VIOLATED",
        })

    decision = final_decision(comparison_results)
    next_task = (
        "R18H_GF5_PARTITION_OBJECTIVE_CORRECTION"
        if decision == "R18G_B2_ADR_PARTITION_OBJECTIVE_SUPPORTED"
        else "R18H_FREEZE_CURRENT_GF5_PARTITION_POLICY"
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    reference_output = output_dir / "r18g_targeted_partition_reference.csv"
    candidate_output = output_dir / "r18g_targeted_partition_candidate_evaluation.csv"
    summary_output = output_dir / "r18g_targeted_partition_summary.json"
    for path, fields, rows in (
        (reference_output, REFERENCE_FIELDS, reference_rows),
        (candidate_output, CANDIDATE_FIELDS, candidate_rows),
    ):
        with path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
            writer.writeheader()
            writer.writerows(rows)
    summary = {
        "status": decision,
        "next": next_task,
        "next_later": "R18G_C_SEALED_HOLDOUT_VALIDATION",
        "submission": {
            "bytes": validated["submission_bytes"],
            "sha256": validated["submission_sha256"],
        },
        "targeted_reference": "VALIDATED_AND_FROZEN",
        "human_evidence_strength": HUMAN_EVIDENCE_STRENGTH,
        "comparison_results": comparison_results,
        "human_co_membership": co_membership,
        "development_whole_group_counts_unchanged": {
            "SAME_ONE_IDENTITY": 11, "NOT_ONE_IDENTITY": 32,
            "INSUFFICIENT_INFORMATION": 5,
        },
        "source_fingerprints": source,
        "runtime_detector_or_gf5_changes": 0,
        "provider_calls": 0,
        "sealed_holdout": "UNSEEN",
        "holdout_validation_contract": [
            "whole-group SAME / NOT_ONE / INSUFFICIENT",
            "human-mixed accepted groups", "Likely vs Review behavior",
            "partition-sensitive holdout cases", "cannot-link safety",
            "deferred-family behavior", "cross-site behavior",
        ],
    }
    write_manifest(summary_output, summary)
    summary["artifacts"] = {
        path.name: {"bytes": path.stat().st_size, "sha256": sha256_file(path)}
        for path in (reference_output, candidate_output, summary_output)
    }
    return summary


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--submission", required=True, type=Path)
    parser.add_argument("--development", required=True, type=Path)
    parser.add_argument("--development-reference", required=True, type=Path)
    parser.add_argument("--source-mapping", required=True, type=Path)
    parser.add_argument("--targeted-mapping", required=True, type=Path)
    parser.add_argument("--database", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    result = evaluate(
        args.submission, args.development, args.development_reference,
        args.source_mapping, args.targeted_mapping, args.database, args.output_dir,
    )
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
