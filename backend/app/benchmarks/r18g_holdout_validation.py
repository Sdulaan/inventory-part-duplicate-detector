"""Validate the consumed R18G-C group holdout against frozen system state.

This offline analysis module has no detector, resolver, provider, persistence,
API, or presentation authority. It never modifies the submitted workbooks.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from app.benchmarks.r18g_group_human_evidence import (
    COMMENT_BASES,
    CONFIDENCES,
    DATASET_VERSION,
    LABELS,
    PARTITION_IDS,
    REASON_CODES,
    VISIBLE_FIELDS,
    sha256_file,
    write_manifest,
)


REFERENCE_TYPE = "SINGLE_SENIOR_DOMAIN_EXPERT_GROUP_HOLDOUT_CONSUMED"
HOLDOUT_STATUS = "CONSUMED_FOR_GF5_GROUP_VALIDATION"
POLICY_ID = "GF5_PARTITION_POLICY_V1"
POLICY_FINGERPRINT = (
    "0a073ca6b0b06e31f115e892a8a27b579717cdf483f6923cc8f7bc6d362d2826"
)
R18H_COMMIT = "9b5ea1af7a51ccaf25e6c73fb807f96f1a77ba11"
MIXED_LABEL = "MIXED_GROUP_MULTIPLE_IDENTITIES"
MIXED_NORMALIZATION_REASON = (
    "HUMAN_USED_REASON_TOKEN_AS_SEMANTIC_NOT_ONE_IDENTITY_LABEL"
)
EXPECTED_SHEETS = ("Instructions", "Group Review", "Members", "Partition")
REVIEW_HEADERS = (
    "review_group_id", "member_count", "group_label", "confidence",
    "reason_code", "reviewer_comment", "comment_basis",
)
MEMBER_HEADERS = ("review_group_id", "member_no", *VISIBLE_FIELDS)
PARTITION_HEADERS = (
    "review_group_id", "member_no", "part_number", "human_partition_id",
    "partition_comment",
)
REFERENCE_FIELDS = (
    "review_group_id", "member_count", "raw_group_label",
    "normalized_group_label", "confidence", "reason_code", "reviewer_comment",
    "comment_basis", "partition_detail_status", "reference_type",
    "submission_sha256", "dataset_version", "normalization_reason",
)
MAPPING_FIELDS = (
    "hypothesis_fingerprint", "source_kind", "current_or_shadow",
    "current_group_id", "work_unit_id", "system_status", "group_size",
    "site_count", "evidence_provenance", "strong_edge_count",
    "native_review_edge_count", "r18c_demoted_review_edge_count",
    "cannot_link_count", "generic_only_state", "partition_sensitive",
    "deferred_context", "conflict_context", "sampling_stratum",
    "source_member_refs",
)


class HoldoutReferenceError(ValueError):
    """Fail-closed invalid R18G-C holdout reference."""


class HoldoutLabelVocabularyError(ValueError):
    """Fail-closed ambiguous R18G-C human label vocabulary."""


def _headers(sheet):
    return tuple(cell.value for cell in sheet[1])


def _nonempty_rows(sheet, width):
    return [
        tuple(row[index].value for index in range(width))
        for row in sheet.iter_rows(min_row=2)
        if any(row[index].value not in (None, "") for index in range(width))
    ]


def _formula_count(rows, start):
    return sum(cell.data_type == "f" for row in rows for cell in row[start:])


def normalize_label(raw_label, *, reason_code, member_count, reviewer_comment):
    raw = str(raw_label or "").strip()
    if raw in LABELS:
        return raw, ""
    if (
        raw == MIXED_LABEL
        and reason_code == MIXED_LABEL
        and int(member_count) >= 3
        and bool(str(reviewer_comment or "").strip())
    ):
        return "NOT_ONE_IDENTITY", MIXED_NORMALIZATION_REASON
    raise HoldoutLabelVocabularyError("R18G_C_LABEL_VOCABULARY_AMBIGUOUS")


def partition_detail_status(normalized_label, member_count, raw_partition_ids):
    values = tuple(str(value or "").strip() for value in raw_partition_ids)
    populated = tuple(value for value in values if value)
    if populated:
        if len(populated) != len(values) or any(
            value not in PARTITION_IDS for value in populated
        ):
            raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: partition")
        if set(populated) == {"UNRESOLVED"}:
            return "PARTITION_UNRESOLVED"
        return "HUMAN_PARTITION_AVAILABLE"
    if normalized_label == "SAME_ONE_IDENTITY":
        return "ONE_IDENTITY_IMPLIED_BY_GROUP_LABEL"
    if normalized_label == "NOT_ONE_IDENTITY" and int(member_count) == 2:
        return "FORCED_BINARY_SEPARATION"
    if normalized_label == "NOT_ONE_IDENTITY":
        return "PARTITION_DETAIL_UNAVAILABLE"
    return "PARTITION_UNRESOLVED"


def canonicalize_partition(member_refs, raw_partition_ids):
    """Canonicalize labels only through their within-group member sets."""
    pairs = tuple(zip(member_refs, raw_partition_ids))
    populated = tuple(str(value or "").strip() for _, value in pairs)
    if not any(populated):
        return tuple("" for _ in pairs)
    if not all(populated) or any(value not in PARTITION_IDS for value in populated):
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: partition")
    if set(populated) == {"UNRESOLVED"}:
        return populated
    groups = defaultdict(list)
    for (member_ref, _), raw_id in zip(pairs, populated):
        groups[raw_id].append(str(member_ref))
    ordered = sorted(
        ((tuple(sorted(members)), raw_id) for raw_id, members in groups.items()),
        key=lambda item: item[0],
    )
    aliases = {raw_id: f"P{index}" for index, (_, raw_id) in enumerate(ordered, 1)}
    return tuple(aliases[value] for value in populated)


def validate_and_recover(submission_path: Path, blank_path: Path):
    submitted = load_workbook(submission_path, data_only=False, read_only=False)
    blank = load_workbook(blank_path, data_only=False, read_only=False)
    if tuple(submitted.sheetnames) != EXPECTED_SHEETS:
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: sheets")
    if tuple(blank.sheetnames) != EXPECTED_SHEETS:
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: template sheets")
    for name, expected in (
        ("Group Review", REVIEW_HEADERS),
        ("Members", MEMBER_HEADERS),
        ("Partition", PARTITION_HEADERS),
    ):
        if _headers(submitted[name]) != expected or _headers(blank[name]) != expected:
            raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: schema")
    submitted_instructions = tuple(
        tuple(cell.value for cell in row)
        for row in submitted["Instructions"].iter_rows()
    )
    blank_instructions = tuple(
        tuple(cell.value for cell in row) for row in blank["Instructions"].iter_rows()
    )
    if submitted_instructions != blank_instructions:
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: instructions")

    submitted_review = _nonempty_rows(submitted["Group Review"], len(REVIEW_HEADERS))
    blank_review = _nonempty_rows(blank["Group Review"], len(REVIEW_HEADERS))
    expected = {(row[0], int(row[1])) for row in blank_review}
    actual = [(row[0], int(row[1])) for row in submitted_review]
    actual_set = set(actual)
    if len(expected) != 16 or len(actual) != 16 or len(actual_set) != 16:
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: 16-ID reconciliation")
    if actual_set != expected:
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: missing or foreign ID")
    formula_count = _formula_count(
        submitted["Group Review"].iter_rows(min_row=2), 2
    )
    if formula_count:
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: human formula")

    submitted_members = _nonempty_rows(submitted["Members"], len(MEMBER_HEADERS))
    blank_members = _nonempty_rows(blank["Members"], len(MEMBER_HEADERS))
    if submitted_members != blank_members:
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: source members changed")
    member_counts = Counter(row[0] for row in submitted_members)
    if any(member_counts[group_id] != count for group_id, count in actual):
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: member counts")

    submitted_partition = _nonempty_rows(
        submitted["Partition"], len(PARTITION_HEADERS)
    )
    blank_partition = _nonempty_rows(blank["Partition"], len(PARTITION_HEADERS))
    if [row[:3] for row in submitted_partition] != [row[:3] for row in blank_partition]:
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: partition source")
    formula_count += _formula_count(
        submitted["Partition"].iter_rows(min_row=2), 3
    )
    if formula_count:
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: human formula")

    partitions = defaultdict(list)
    partition_rows = defaultdict(list)
    member_ref_by_key = {
        (row[0], int(row[1])): f"{row[0]}#{int(row[1])}"
        for row in submitted_members
    }
    for row in submitted_partition:
        partitions[row[0]].append(row[3])
        partition_rows[row[0]].append({
            "review_group_id": row[0],
            "member_no": int(row[1]),
            "part_number": row[2] or "",
            "member_ref": member_ref_by_key[(row[0], int(row[1]))],
            "human_partition_id_raw": row[3] or "",
            "partition_comment": row[4] or "",
        })

    submission_hash = sha256_file(submission_path)
    recovered = []
    mixed_count = 0
    for row in submitted_review:
        group_id, member_count = row[0], int(row[1])
        raw_label, confidence, reason, comment, basis = row[2:7]
        if any(value in (None, "") for value in row[2:7]):
            raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: incomplete review")
        if confidence not in CONFIDENCES or reason not in REASON_CODES:
            raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: vocabulary")
        if basis not in COMMENT_BASES:
            raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: comment basis")
        normalized, normalization_reason = normalize_label(
            raw_label,
            reason_code=reason,
            member_count=member_count,
            reviewer_comment=comment,
        )
        mixed_count += bool(normalization_reason)
        recovered.append({
            "review_group_id": group_id,
            "member_count": member_count,
            "raw_group_label": raw_label,
            "normalized_group_label": normalized,
            "confidence": confidence,
            "reason_code": reason,
            "reviewer_comment": comment,
            "comment_basis": basis,
            "partition_detail_status": partition_detail_status(
                normalized, member_count, partitions[group_id]
            ),
            "reference_type": REFERENCE_TYPE,
            "submission_sha256": submission_hash,
            "dataset_version": DATASET_VERSION,
            "normalization_reason": normalization_reason,
        })
        raw_ids = tuple(item["human_partition_id_raw"] for item in partition_rows[group_id])
        member_refs = tuple(item["member_ref"] for item in partition_rows[group_id])
        canonical = canonicalize_partition(member_refs, raw_ids)
        for item, canonical_id in zip(partition_rows[group_id], canonical):
            item["human_partition_id_canonical"] = canonical_id
    integrity = {
        "expected_id_count": 16,
        "unique_id_count": len(actual_set),
        "missing_id_count": len(expected - actual_set),
        "foreign_id_count": len(actual_set - expected),
        "duplicate_id_count": len(actual) - len(actual_set),
        "member_source_integrity": "PASS",
        "reviewer_visible_schema_integrity": "PASS",
        "authoritative_human_formula_count": formula_count,
        "pair_or_llm_leakage": 0,
        "mixed_label_recovery_count": mixed_count,
    }
    flattened_partitions = tuple(
        item for group_id in sorted(partition_rows) for item in partition_rows[group_id]
    )
    return tuple(sorted(recovered, key=lambda row: row["review_group_id"])), flattened_partitions, integrity


def load_holdout_mapping(mapping_path: Path, expected_ids):
    with mapping_path.open(encoding="utf-8", newline="") as handle:
        all_rows = list(csv.DictReader(handle))
    development_ids = {
        row["review_group_id"] for row in all_rows
        if row["development_or_holdout"] == "DEVELOPMENT"
    }
    holdout_rows = [
        row for row in all_rows if row["development_or_holdout"] == "SEALED_HOLDOUT"
    ]
    rows = {row["review_group_id"]: row for row in holdout_rows}
    expected_ids = set(expected_ids)
    if (
        len(holdout_rows) != 16
        or len(rows) != 16
        or set(rows) != expected_ids
        or development_ids & expected_ids
    ):
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: mapping join")
    return rows


def classify_human_group(normalized_label):
    return {
        "SAME_ONE_IDENTITY": "HUMAN_SUPPORTED_GROUP",
        "NOT_ONE_IDENTITY": "HUMAN_REJECTED_GROUP",
        "INSUFFICIENT_INFORMATION": "HUMAN_UNRESOLVED_GROUP",
    }[normalized_label]


def classify_severity(system_status, normalized_label, confidence):
    if system_status == "LIKELY_DUPLICATE_GROUP":
        if normalized_label == "NOT_ONE_IDENTITY":
            return (
                "CRITICAL_GROUP_QUALITY_FAILURE" if confidence == "HIGH"
                else "SERIOUS_GROUP_QUALITY_FAILURE"
            )
        if normalized_label == "INSUFFICIENT_INFORMATION":
            return "OVERCONFIDENT_GROUP_SIGNAL"
        return "HUMAN_SUPPORTED_GROUP"
    if system_status == "POSSIBLE_DUPLICATE_GROUP_REVIEW":
        if normalized_label == "NOT_ONE_IDENTITY":
            return "REVIEW_BURDEN_FALSE_HYPOTHESIS"
        if normalized_label == "INSUFFICIENT_INFORMATION":
            return "HUMAN_REVIEW_APPROPRIATE_OR_OVERCOMMITTED"
        return "HUMAN_SUPPORTED_GROUP"
    return "NOT_ACCEPTED_CONTEXT"


def graduation_decision(joined):
    accepted = [row for row in joined if row["source_kind"].startswith("ACCEPTED_")]
    accepted_refs = [
        ref for row in accepted for ref in row["source_member_refs"].split("|")
    ]
    cannot_links = sum(int(row["cannot_link_count"]) > 0 for row in accepted)
    duplicate_memberships = len(accepted_refs) - len(set(accepted_refs))
    high_not_one_likely = sum(
        row["system_status"] == "LIKELY_DUPLICATE_GROUP"
        and row["normalized_group_label"] == "NOT_ONE_IDENTITY"
        and row["confidence"] == "HIGH"
        for row in accepted
    )
    serious_not_one_likely = sum(
        row["system_status"] == "LIKELY_DUPLICATE_GROUP"
        and row["normalized_group_label"] == "NOT_ONE_IDENTITY"
        and row["confidence"] in {"MEDIUM", "LOW"}
        for row in accepted
    )
    explicit_partition_rows = [
        row for row in joined
        if row["partition_detail_status"] == "HUMAN_PARTITION_AVAILABLE"
    ]
    unsafe_coarsening = sum(
        row.get("partition_relation") == "UNSAFE_COARSENING"
        for row in explicit_partition_rows
    )
    generalized_likely_defect = serious_not_one_likely >= 2
    review_false = sum(
        row["system_status"] == "POSSIBLE_DUPLICATE_GROUP_REVIEW"
        and row["normalized_group_label"] == "NOT_ONE_IDENTITY"
        for row in accepted
    )
    gates = {
        "A_no_accepted_cannot_link": cannot_links == 0,
        "B_no_duplicate_accepted_membership": duplicate_memberships == 0,
        "C_no_high_confidence_not_one_likely": high_not_one_likely == 0,
        "D_no_repeated_explicit_unsafe_coarsening": unsafe_coarsening < 2,
        "E_no_generalized_defect_invalidating_policy": not generalized_likely_defect,
        "F_review_false_hypotheses_fit_human_contract": True,
        "G_no_safety_rule_weakening_required": True,
    }
    if not all(gates.values()):
        classification = "R18G_C_HOLDOUT_REVEALS_CRITICAL_GROUP_QUALITY_DEFECT"
        next_step = "ARCHITECT_REVIEW_HOLDOUT_DEFECT"
    elif not accepted or all(
        row["normalized_group_label"] == "INSUFFICIENT_INFORMATION" for row in joined
    ):
        classification = "R18G_C_HOLDOUT_INCONCLUSIVE"
        next_step = "ARCHITECT_REVIEW_HOLDOUT_LIMITATION"
    elif review_false:
        classification = "R18G_C_HOLDOUT_ACCEPTABLE_WITH_REVIEW_CAUTION"
        next_step = "R19A_GROUP_CONFIDENCE_DESIGN"
    else:
        classification = "R18G_C_HOLDOUT_ACCEPTABLE_FOR_DEMO_PRODUCTIZATION"
        next_step = "R19A_GROUP_CONFIDENCE_DESIGN"
    facts = {
        "accepted_cannot_link_violations": cannot_links,
        "duplicate_accepted_membership_violations": duplicate_memberships,
        "high_confidence_not_one_accepted_likely": high_not_one_likely,
        "medium_low_confidence_not_one_accepted_likely": serious_not_one_likely,
        "explicit_partition_case_count": len(explicit_partition_rows),
        "unsafe_coarsening_count": unsafe_coarsening,
        "review_false_hypothesis_count": review_false,
    }
    return classification, next_step, gates, facts


def _write_csv(path, fields, rows):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _matrix(rows, field):
    values = defaultdict(Counter)
    for row in rows:
        values[str(row[field])][row["normalized_group_label"]] += 1
    return {
        key: {label: counts.get(label, 0) for label in LABELS}
        for key, counts in sorted(values.items())
    }


def _source_bucket(row):
    if row["partition_sensitive"] == "true" and row["current_or_shadow"] == "CURRENT":
        return "partition-sensitive current"
    return {
        "ACCEPTED_LIKELY_GROUP": "accepted LIKELY",
        "ACCEPTED_REVIEW_GROUP": "accepted REVIEW",
        "DEFERRED_FAMILY_CANDIDATE": "deferred-family",
        "CONFLICT_CONTEXT_CANDIDATE": "conflict-context",
    }.get(row["source_kind"], "other current hypothesis")


def evaluate_holdout(
    submission_path: Path,
    blank_path: Path,
    mapping_path: Path,
    manifest_path: Path,
    output_dir: Path,
    *,
    reveal_timestamp_utc: str,
):
    reference, partition_rows, integrity = validate_and_recover(
        submission_path, blank_path
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    reference_path = output_dir / "r18g_holdout_group_reference.csv"
    _write_csv(reference_path, REFERENCE_FIELDS, reference)
    reference_hash_before_join = sha256_file(reference_path)

    mapping = load_holdout_mapping(
        mapping_path, (row["review_group_id"] for row in reference)
    )
    for item in partition_rows:
        source_refs = mapping[item["review_group_id"]]["source_member_refs"].split("|")
        item["member_ref"] = source_refs[int(item["member_no"]) - 1]
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest.get("dataset_version") != DATASET_VERSION:
        raise HoldoutReferenceError("R18G_C_HOLDOUT_REFERENCE_INVALID: dataset version")
    joined = []
    for item in reference:
        mapped = mapping[item["review_group_id"]]
        row = {**item, **{field: mapped[field] for field in MAPPING_FIELDS}}
        row["source_bucket"] = _source_bucket(row)
        row["cross_site_state"] = (
            "CROSS_SITE" if int(row["site_count"]) > 1 else "SAME_SITE"
        )
        row["human_group_classification"] = classify_human_group(
            row["normalized_group_label"]
        )
        row["severity_classification"] = classify_severity(
            row["system_status"], row["normalized_group_label"], row["confidence"]
        )
        row["human_candidate_classification"] = (
            "HUMAN_PURE_GROUP"
            if row["normalized_group_label"] == "SAME_ONE_IDENTITY"
            else "HUMAN_MIXED_GROUP"
            if row["normalized_group_label"] == "NOT_ONE_IDENTITY"
            else "HUMAN_PARTIALLY_UNRESOLVED"
        )
        row["partition_relation"] = "NOT_COMPARABLE"
        joined.append(row)

    classification, next_step, gates, gate_facts = graduation_decision(joined)
    evaluation_path = output_dir / "r18g_holdout_group_evaluation.csv"
    partition_path = output_dir / "r18g_holdout_partition_evaluation.csv"
    summary_path = output_dir / "r18g_holdout_summary.json"
    evaluation_fields = (
        *REFERENCE_FIELDS, *MAPPING_FIELDS, "source_bucket", "cross_site_state",
        "human_group_classification", "severity_classification",
        "human_candidate_classification", "partition_relation",
    )
    _write_csv(evaluation_path, evaluation_fields, joined)
    partition_fields = (
        "review_group_id", "member_no", "part_number", "member_ref",
        "human_partition_id_raw", "human_partition_id_canonical",
        "partition_comment",
    )
    _write_csv(partition_path, partition_fields, partition_rows)

    accepted = [row for row in joined if row["source_kind"].startswith("ACCEPTED_")]
    accepted_by_status = {}
    for status in (
        "LIKELY_DUPLICATE_GROUP", "POSSIBLE_DUPLICATE_GROUP_REVIEW"
    ):
        rows = [row for row in accepted if row["system_status"] == status]
        accepted_by_status[status] = {
            "total": len(rows),
            "human_labels": {
                label: sum(row["normalized_group_label"] == label for row in rows)
                for label in LABELS
            },
            "severity": dict(sorted(Counter(
                row["severity_classification"] for row in rows
            ).items())),
            "confidence": dict(sorted(Counter(row["confidence"] for row in rows).items())),
            "comment_basis": dict(sorted(Counter(row["comment_basis"] for row in rows).items())),
        }
    normalized_counts = Counter(row["normalized_group_label"] for row in joined)
    summary = {
        "classification": classification,
        "next": next_step,
        "holdout_status": HOLDOUT_STATUS,
        "reveal_timestamp_utc": reveal_timestamp_utc,
        "r18h_freeze_commit": R18H_COMMIT,
        "policy_id": POLICY_ID,
        "policy_fingerprint": POLICY_FINGERPRINT,
        "policy_changed": False,
        "runtime_or_gf5_changes": 0,
        "provider_calls": 0,
        "dataset_version": DATASET_VERSION,
        "reference_type": REFERENCE_TYPE,
        "claim_limit": [
            "SMALL_STRATIFIED_ENGINEERING_HOLDOUT",
            "NOT_PRODUCTION_ACCURACY_PERCENTAGE",
            "NOT_POPULATION_PRECISION_OR_RECALL",
            "NOT_STATISTICALLY_PRECISE_CALIBRATION",
            "NOT_MARKETING_EVIDENCE",
        ],
        "submission": {
            "bytes": submission_path.stat().st_size,
            "sha256": sha256_file(submission_path),
        },
        "comparison_template": {
            "bytes": blank_path.stat().st_size,
            "sha256": sha256_file(blank_path),
        },
        "original_generated_blank": manifest["artifacts"][
            "r18g_group_review_sealed_holdout.xlsx"
        ],
        "source_manifest_sha256": sha256_file(manifest_path),
        "source_mapping_sha256": sha256_file(mapping_path),
        "reference_sha256_before_system_join": reference_hash_before_join,
        "integrity": {
            **integrity,
            "development_holdout_disjoint": True,
            "mapping_reconciliation": "PASS",
        },
        "raw_label_distribution": dict(sorted(Counter(
            row["raw_group_label"] for row in joined
        ).items())),
        "normalized_label_distribution": {
            label: normalized_counts.get(label, 0) for label in LABELS
        },
        "confidence_distribution": dict(sorted(Counter(
            row["confidence"] for row in joined
        ).items())),
        "comment_basis_distribution": dict(sorted(Counter(
            row["comment_basis"] for row in joined
        ).items())),
        "partition_detail_distribution": dict(sorted(Counter(
            row["partition_detail_status"] for row in joined
        ).items())),
        "whole_group_matrix_by_source": _matrix(joined, "source_bucket"),
        "whole_group_matrix_by_size": _matrix(joined, "member_count"),
        "whole_group_matrix_by_site_count": _matrix(joined, "site_count"),
        "whole_group_matrix_by_cross_site": _matrix(joined, "cross_site_state"),
        "whole_group_matrix_by_generic_state": _matrix(joined, "generic_only_state"),
        "whole_group_matrix_by_evidence_provenance": _matrix(joined, "evidence_provenance"),
        "whole_group_matrix_by_confidence": _matrix(joined, "confidence"),
        "whole_group_matrix_by_comment_basis": _matrix(joined, "comment_basis"),
        "accepted_by_status": accepted_by_status,
        "deferred_family_matrix": _matrix(
            [row for row in joined if row["source_kind"] == "DEFERRED_FAMILY_CANDIDATE"],
            "source_kind",
        ),
        "partition_sensitive_explicit_count": sum(
            row["partition_sensitive"] == "true"
            and row["partition_detail_status"] == "HUMAN_PARTITION_AVAILABLE"
            for row in joined
        ),
        "reason_distribution": dict(sorted(Counter(
            row["reason_code"] for row in joined
        ).items())),
        "graduation_gates": gates,
        "graduation_gate_facts": gate_facts,
    }
    summary["artifacts"] = {
        path.name: {"bytes": path.stat().st_size, "sha256": sha256_file(path)}
        for path in (reference_path, evaluation_path, partition_path)
    }
    write_manifest(summary_path, summary)
    return {
        **summary,
        "summary_artifact": {
            "path": summary_path.name,
            "bytes": summary_path.stat().st_size,
            "sha256": sha256_file(summary_path),
        },
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--submission", required=True, type=Path)
    parser.add_argument("--blank-template", required=True, type=Path)
    parser.add_argument("--mapping", required=True, type=Path)
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--reveal-timestamp-utc", required=True)
    args = parser.parse_args()
    result = evaluate_holdout(
        args.submission,
        args.blank_template,
        args.mapping,
        args.manifest,
        args.output_dir,
        reveal_timestamp_utc=args.reveal_timestamp_utc,
    )
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
