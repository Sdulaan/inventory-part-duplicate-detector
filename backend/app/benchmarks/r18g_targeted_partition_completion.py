"""Prepare the two-group blinded R18G-B1 partition-completion workbook."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Protection
from openpyxl.worksheet.datavalidation import DataValidation

from app.benchmarks.r18g_development_group_evaluation import (
    REFERENCE_TYPE,
    assert_holdout_firewall,
)
from app.benchmarks.r18g_group_human_evidence import (
    PARTITION_IDS,
    VISIBLE_FIELDS,
    _new_workbook,
    _safe_cell,
    _save_reproducible,
    sha256_file,
)


TARGET_IDS = (
    "R18G-D09438A953125080",
    "R18G-FFF38D9F3046F48C",
)
EXPECTED_MEMBER_COUNTS = {
    "R18G-D09438A953125080": 5,
    "R18G-FFF38D9F3046F48C": 4,
}
GROUP_HEADERS = (
    "review_group_id", "member_count", "previous_raw_group_label",
    "previous_normalized_group_label", "previous_confidence",
    "previous_reason_code", "previous_reviewer_comment",
    "previous_comment_basis",
)
MEMBER_HEADERS = ("review_group_id", "member_no", *VISIBLE_FIELDS)
PARTITION_HEADERS = (
    "review_group_id", "member_no", "part_number", "human_partition_id",
    "partition_comment",
)
INTERNAL_MAPPING_FIELDS = (
    "review_group_id", "hypothesis_fingerprint", "comparison_set_id",
    "current_or_shadow", "work_unit_id", "source_member_refs", "member_count",
)
FORBIDDEN_HEADERS = {
    "current_or_shadow", "current_group_id", "work_unit_id", "source_kind",
    "system_status", "strong_edge_count", "native_review_edge_count",
    "r18c_demoted_review_edge_count", "cannot_link_count", "score",
    "pair_label", "reviewer_b", "llm",
}


class TargetedPartitionPackageError(ValueError):
    """Fail-closed invalid targeted package or source evidence."""


def comparison_set_id(work_unit_id, source_member_refs):
    payload = f"r18g-b1|{work_unit_id}|{source_member_refs}"
    return "R18G-CMP-" + hashlib.sha256(payload.encode()).hexdigest()[:16].upper()


def _csv_rows(path):
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def load_target_sources(submission_path: Path, reference_path: Path,
                        mapping_path: Path):
    assert_holdout_firewall(submission_path, reference_path, mapping_path)
    reference_rows = {
        row["review_group_id"]: row for row in _csv_rows(reference_path)
        if row["review_group_id"] in TARGET_IDS
    }
    if set(reference_rows) != set(TARGET_IDS):
        raise TargetedPartitionPackageError(
            "R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: reference IDs"
        )
    if any(row["reference_type"] != REFERENCE_TYPE for row in reference_rows.values()):
        raise TargetedPartitionPackageError(
            "R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: reference type"
        )
    mapping_rows = {
        row["review_group_id"]: row for row in _csv_rows(mapping_path)
        if row["review_group_id"] in TARGET_IDS
        and row["development_or_holdout"] == "DEVELOPMENT"
    }
    if set(mapping_rows) != set(TARGET_IDS):
        raise TargetedPartitionPackageError(
            "R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: mapping IDs"
        )

    workbook = load_workbook(submission_path, data_only=False, read_only=True)
    review_headers = tuple(cell.value for cell in workbook["Group Review"][1])
    review_rows = {
        row[0]: row for row in workbook["Group Review"].iter_rows(
            min_row=2, values_only=True
        ) if row[0] in TARGET_IDS
    }
    if set(review_rows) != set(TARGET_IDS):
        raise TargetedPartitionPackageError(
            "R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: submission IDs"
        )
    review_index = {name: review_headers.index(name) for name in review_headers}
    for group_id, reference in reference_rows.items():
        submitted = review_rows[group_id]
        comparisons = {
            "member_count": "member_count",
            "raw_group_label": "group_label",
            "confidence": "confidence",
            "reason_code": "reason_code",
            "reviewer_comment": "reviewer_comment",
            "comment_basis": "comment_basis",
        }
        if any(
            str(reference[reference_field]) != str(submitted[review_index[submitted_field]])
            for reference_field, submitted_field in comparisons.items()
        ):
            raise TargetedPartitionPackageError(
                "R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: prior judgment mismatch"
            )

    members = []
    for row in workbook["Members"].iter_rows(min_row=2, values_only=True):
        if row[0] in TARGET_IDS:
            members.append(tuple(row[:len(MEMBER_HEADERS)]))
    counts = {}
    for group_id in TARGET_IDS:
        group_members = [row for row in members if row[0] == group_id]
        counts[group_id] = len(group_members)
        expected = EXPECTED_MEMBER_COUNTS[group_id]
        if len(group_members) != expected or {
            int(row[1]) for row in group_members
        } != set(range(1, expected + 1)):
            raise TargetedPartitionPackageError(
                "R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: member reconciliation"
            )
        if int(reference_rows[group_id]["member_count"]) != expected:
            raise TargetedPartitionPackageError(
                "R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: reference member count"
            )
        if len(mapping_rows[group_id]["source_member_refs"].split("|")) != expected:
            raise TargetedPartitionPackageError(
                "R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: member-set mapping"
            )
    return reference_rows, mapping_rows, tuple(sorted(members, key=lambda row: (row[0], int(row[1]))))


def _protect_read_only(sheet):
    sheet.protection.sheet = True
    sheet.protection.selectLockedCells = False
    sheet.protection.selectUnlockedCells = False


def write_targeted_workbook(path: Path, reference_rows, members):
    workbook = _new_workbook("R18G targeted partition completion")
    instructions = workbook.active
    instructions.title = "Instructions"
    instruction_lines = (
        "This workbook contains only two groups that you previously reviewed.",
        "The earlier whole-group judgments remain unchanged and are not being reconsidered.",
        "For each member, assign P1, P2, P3, and so on in the Partition sheet.",
        "Members with the same P-number represent the same underlying physical/business inventory identity.",
        "Different P-numbers represent distinct identities.",
        "Use UNRESOLVED when you cannot safely determine a member's partition.",
        "Different sites or ERP part numbers alone do not prove different identity.",
        "Use professional domain judgment and do not guess.",
        "Only human_partition_id and partition_comment cells are editable.",
    )
    instructions.append(["R18G targeted partition completion"])
    for line in instruction_lines:
        instructions.append([line])
    instructions.column_dimensions["A"].width = 120

    groups = workbook.create_sheet("Groups")
    groups.append(GROUP_HEADERS)
    for group_id in TARGET_IDS:
        row = reference_rows[group_id]
        groups.append([
            group_id, int(row["member_count"]), row["raw_group_label"],
            row["normalized_group_label"], row["confidence"], row["reason_code"],
            _safe_cell(row["reviewer_comment"]), row["comment_basis"],
        ])

    member_sheet = workbook.create_sheet("Members")
    member_sheet.append(MEMBER_HEADERS)
    partition = workbook.create_sheet("Partition")
    partition.append(PARTITION_HEADERS)
    for row in members:
        member_sheet.append([_safe_cell(value) for value in row])
        partition.append([row[0], int(row[1]), _safe_cell(row[2]), "", ""])
    validation = DataValidation(
        type="list", formula1='"' + ",".join(PARTITION_IDS) + '"', allow_blank=True
    )
    partition.add_data_validation(validation)
    validation.add(f"D2:D{partition.max_row}")

    for sheet in (groups, member_sheet, partition):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(70, max(14, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    for row in partition.iter_rows(min_row=2, min_col=4, max_col=5):
        for cell in row:
            cell.protection = Protection(locked=False)
    for sheet in (instructions, groups, member_sheet, partition):
        _protect_read_only(sheet)
    partition.protection.selectUnlockedCells = True
    _save_reproducible(workbook, path)


def write_internal_mapping(path: Path, mapping_rows):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(
            handle, fieldnames=INTERNAL_MAPPING_FIELDS, lineterminator="\n"
        )
        writer.writeheader()
        for group_id in TARGET_IDS:
            source = mapping_rows[group_id]
            writer.writerow({
                "review_group_id": group_id,
                "hypothesis_fingerprint": source["hypothesis_fingerprint"],
                "comparison_set_id": comparison_set_id(
                    source["work_unit_id"], source["source_member_refs"]
                ),
                "current_or_shadow": source["current_or_shadow"],
                "work_unit_id": source["work_unit_id"],
                "source_member_refs": source["source_member_refs"],
                "member_count": source["group_size"],
            })


def validate_targeted_package(workbook_path: Path, mapping_path: Path,
                              reference_rows, source_members):
    assert_holdout_firewall(workbook_path, mapping_path)
    workbook = load_workbook(workbook_path, data_only=False, read_only=False)
    if workbook.sheetnames != ["Instructions", "Groups", "Members", "Partition"]:
        raise TargetedPartitionPackageError("R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: sheets")
    groups = workbook["Groups"]
    if tuple(cell.value for cell in groups[1]) != GROUP_HEADERS:
        raise TargetedPartitionPackageError("R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: group headers")
    group_rows = list(groups.iter_rows(min_row=2, values_only=True))
    if tuple(row[0] for row in group_rows) != TARGET_IDS:
        raise TargetedPartitionPackageError("R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: exact IDs")
    for row in group_rows:
        reference = reference_rows[row[0]]
        expected = (
            row[0], int(reference["member_count"]), reference["raw_group_label"],
            reference["normalized_group_label"], reference["confidence"],
            reference["reason_code"], reference["reviewer_comment"],
            reference["comment_basis"],
        )
        if row != expected:
            raise TargetedPartitionPackageError("R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: judgment")
    member_rows = tuple(
        tuple(row) for row in workbook["Members"].iter_rows(min_row=2, values_only=True)
    )
    if member_rows != source_members:
        raise TargetedPartitionPackageError("R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: source fields")
    partition_rows = list(workbook["Partition"].iter_rows(min_row=2))
    expected_keys = {(row[0], int(row[1]), row[2]) for row in source_members}
    actual_keys = {(row[0].value, int(row[1].value), row[2].value) for row in partition_rows}
    if actual_keys != expected_keys:
        raise TargetedPartitionPackageError("R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: partition members")
    if any(cell.value not in (None, "") or cell.data_type == "f"
           for row in partition_rows for cell in row[3:5]):
        raise TargetedPartitionPackageError("R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: prefilled partition")
    formulas = {
        item.formula1 for item in workbook["Partition"].data_validations.dataValidation
    }
    if formulas != {'"' + ",".join(PARTITION_IDS) + '"'}:
        raise TargetedPartitionPackageError("R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: dropdown")
    visible_headers = {
        str(cell.value or "").casefold()
        for sheet in workbook for cell in sheet[1]
    }
    if visible_headers & FORBIDDEN_HEADERS:
        raise TargetedPartitionPackageError("R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: leakage")
    mapping = _csv_rows(mapping_path)
    if len(mapping) != 2 or tuple(row["review_group_id"] for row in mapping) != TARGET_IDS:
        raise TargetedPartitionPackageError("R18G_B1_TARGETED_PARTITION_PACKAGE_INVALID: mapping")
    return {
        "target_group_count": 2,
        "target_ids": list(TARGET_IDS),
        "member_counts": EXPECTED_MEMBER_COUNTS,
        "source_field_reconciliation": "PASS",
        "prior_judgment_reconciliation": "PASS",
        "holdout_ids_visible": 0,
        "detector_fields_visible": 0,
        "prefilled_partition_ids": 0,
        "editable_partition_formula_count": 0,
        "partition_dropdown": list(PARTITION_IDS),
    }


def prepare_targeted_package(submission_path: Path, reference_path: Path,
                             source_mapping_path: Path, output_dir: Path):
    assert_holdout_firewall(
        submission_path, reference_path, source_mapping_path, output_dir
    )
    reference, mapping, members = load_target_sources(
        submission_path, reference_path, source_mapping_path
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    workbook_path = output_dir / "r18g_targeted_partition_completion.xlsx"
    mapping_path = output_dir / "r18g_targeted_partition_internal_mapping.csv"
    write_targeted_workbook(workbook_path, reference, members)
    write_internal_mapping(mapping_path, mapping)
    validation = validate_targeted_package(
        workbook_path, mapping_path, reference, members
    )
    return {
        "status": "R18G_B1_AWAITING_TARGETED_PARTITION_HUMAN_INPUT",
        "runtime_or_gf5_changes": 0,
        "provider_calls": 0,
        "sealed_holdout": "UNSEEN",
        "development_group_evidence": "FROZEN",
        "targeted_human_partition_labels": 0,
        "validation": validation,
        "artifacts": {
            path.name: {"bytes": path.stat().st_size, "sha256": sha256_file(path)}
            for path in (workbook_path, mapping_path)
        },
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--development-submission", required=True, type=Path)
    parser.add_argument("--development-reference", required=True, type=Path)
    parser.add_argument("--source-mapping", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    result = prepare_targeted_package(
        args.development_submission, args.development_reference,
        args.source_mapping, args.output_dir,
    )
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
