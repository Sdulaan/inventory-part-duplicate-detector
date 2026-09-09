"""Validate and evaluate only the completed R18G DEVELOPMENT workbook.

The sealed holdout workbook is outside this module's inputs by design.  This
analysis preserves the submitted workbook bytes, writes only ignored evidence
artifacts, and has no provider or runtime authority.
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

from app.benchmarks.r18g_group_human_evidence import (
    COMMENT_BASES,
    CONFIDENCES,
    DATASET_VERSION,
    LABELS,
    PARTITION_IDS,
    REASON_CODES,
    VISIBLE_FIELDS,
    sha256_file,
    write_manifest,
)


REFERENCE_TYPE = "SINGLE_SENIOR_DOMAIN_EXPERT_GROUP_DEVELOPMENT"
MIXED_LABEL = "MIXED_GROUP_MULTIPLE_IDENTITIES"
MIXED_NORMALIZATION_REASON = (
    "HUMAN_USED_REASON_TOKEN_AS_SEMANTIC_NOT_ONE_IDENTITY_LABEL"
)
EXPECTED_SHEETS = ("Instructions", "Group Review", "Members", "Partition")
REVIEW_HEADERS = (
    "review_group_id", "member_count", "group_label", "confidence",
    "reason_code", "reviewer_comment", "comment_basis",
)
MEMBER_HEADERS = ("review_group_id", "member_no", *VISIBLE_FIELDS)
PARTITION_HEADERS = (
    "review_group_id", "member_no", "part_number", "human_partition_id",
    "partition_comment",
)
REFERENCE_FIELDS = (
    "review_group_id", "member_count", "raw_group_label",
    "normalized_group_label", "confidence", "reason_code", "reviewer_comment",
    "comment_basis", "partition_detail_status", "reference_type",
    "submission_sha256", "dataset_version", "normalization_reason",
)
MAPPING_FIELDS = (
    "hypothesis_fingerprint", "source_kind", "current_or_shadow",
    "current_group_id", "work_unit_id", "system_status", "group_size",
    "site_count", "evidence_provenance", "strong_edge_count",
    "native_review_edge_count", "r18c_demoted_review_edge_count",
    "cannot_link_count", "generic_only_state", "partition_sensitive",
    "deferred_context", "conflict_context", "sampling_stratum",
    "source_member_refs",
)


class DevelopmentReferenceError(ValueError):
    """Fail-closed invalid development reference."""


class LabelVocabularyError(ValueError):
    """Fail-closed ambiguous human label vocabulary."""


def assert_holdout_firewall(*paths: Path):
    if any("holdout" in Path(path).name.casefold() for path in paths):
        raise ValueError("R18G_B_HOLDOUT_FIREWALL_BLOCKED")


def _nonempty_rows(sheet, width):
    return [
        tuple(row[index].value for index in range(width))
        for row in sheet.iter_rows(min_row=2)
        if any(row[index].value not in (None, "") for index in range(width))
    ]


def _headers(sheet):
    return tuple(cell.value for cell in sheet[1])


def _formula_count(rows, start):
    return sum(
        cell.data_type == "f" for row in rows for cell in row[start:]
    )


def normalize_label(raw_label, *, reason_code, member_count, reviewer_comment,
                    comment_basis):
    raw = str(raw_label or "").strip()
    if raw in LABELS:
        return raw, ""
    if raw == MIXED_LABEL and (
        reason_code == MIXED_LABEL
        and int(member_count) >= 3
        and bool(str(reviewer_comment or "").strip())
    ):
        return "NOT_ONE_IDENTITY", MIXED_NORMALIZATION_REASON
    raise LabelVocabularyError("R18G_B_LABEL_VOCABULARY_AMBIGUOUS")


def partition_detail_status(normalized_label, member_count, partition_values):
    values = tuple(str(value or "").strip() for value in partition_values)
    populated = tuple(value for value in values if value)
    if populated:
        if len(populated) != len(values):
            raise DevelopmentReferenceError(
                "R18G_B_DEVELOPMENT_REFERENCE_INVALID: partial partition"
            )
        if any(value not in PARTITION_IDS for value in populated):
            raise DevelopmentReferenceError(
                "R18G_B_DEVELOPMENT_REFERENCE_INVALID: partition vocabulary"
            )
        return "HUMAN_PARTITION_AVAILABLE"
    if normalized_label == "SAME_ONE_IDENTITY":
        return "ONE_IDENTITY_IMPLIED_BY_GROUP_LABEL"
    if normalized_label == "NOT_ONE_IDENTITY" and int(member_count) == 2:
        return "FORCED_BINARY_SEPARATION"
    if normalized_label == "NOT_ONE_IDENTITY":
        return "UNAVAILABLE"
    return "UNRESOLVED"


def validate_and_recover(submission_path: Path, blank_path: Path):
    assert_holdout_firewall(submission_path, blank_path)
    submitted = load_workbook(submission_path, data_only=False, read_only=False)
    blank = load_workbook(blank_path, data_only=False, read_only=False)
    if tuple(submitted.sheetnames) != EXPECTED_SHEETS or tuple(blank.sheetnames) != EXPECTED_SHEETS:
        raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: sheets")
    submitted_instructions = tuple(
        tuple(cell.value for cell in row)
        for row in submitted["Instructions"].iter_rows()
    )
    blank_instructions = tuple(
        tuple(cell.value for cell in row)
        for row in blank["Instructions"].iter_rows()
    )
    if submitted_instructions != blank_instructions:
        raise DevelopmentReferenceError(
            "R18G_B_DEVELOPMENT_REFERENCE_INVALID: instructions changed"
        )
    if _headers(submitted["Group Review"]) != REVIEW_HEADERS:
        raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: review headers")
    if _headers(submitted["Members"]) != MEMBER_HEADERS:
        raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: member headers")
    if _headers(submitted["Partition"]) != PARTITION_HEADERS:
        raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: partition headers")

    blank_review = _nonempty_rows(blank["Group Review"], len(REVIEW_HEADERS))
    submitted_review = _nonempty_rows(submitted["Group Review"], len(REVIEW_HEADERS))
    blank_base = {(row[0], int(row[1])) for row in blank_review}
    submitted_base = [(row[0], int(row[1])) for row in submitted_review]
    if (
        len(blank_base) != 48 or len(submitted_base) != 48
        or len(set(submitted_base)) != 48 or set(submitted_base) != blank_base
    ):
        raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: 48-ID reconciliation")
    if _formula_count(submitted["Group Review"].iter_rows(min_row=2), 2):
        raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: human formula")

    blank_members = _nonempty_rows(blank["Members"], len(MEMBER_HEADERS))
    submitted_members = _nonempty_rows(submitted["Members"], len(MEMBER_HEADERS))
    if submitted_members != blank_members:
        raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: source members changed")
    member_counts = Counter(row[0] for row in submitted_members)
    if any(member_counts[group_id] != count for group_id, count in submitted_base):
        raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: member count")

    blank_partition = _nonempty_rows(blank["Partition"], len(PARTITION_HEADERS))
    submitted_partition = _nonempty_rows(submitted["Partition"], len(PARTITION_HEADERS))
    blank_partition_base = [row[:3] for row in blank_partition]
    submitted_partition_base = [row[:3] for row in submitted_partition]
    if submitted_partition_base != blank_partition_base:
        raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: partition source changed")
    if _formula_count(submitted["Partition"].iter_rows(min_row=2), 3):
        raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: partition formula")

    partitions = defaultdict(list)
    for row in submitted_partition:
        partitions[row[0]].append(row[3])
    submission_hash = sha256_file(submission_path)
    recovered = []
    seen = set()
    for row in submitted_review:
        group_id, member_count = row[0], int(row[1])
        if group_id in seen:
            raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: duplicate ID")
        seen.add(group_id)
        raw_label, confidence, reason, comment, basis = row[2:7]
        if any(value in (None, "") for value in (raw_label, confidence, reason, comment, basis)):
            raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: incomplete review")
        if confidence not in CONFIDENCES or reason not in REASON_CODES or basis not in COMMENT_BASES:
            raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: vocabulary")
        normalized, normalization_reason = normalize_label(
            raw_label, reason_code=reason, member_count=member_count,
            reviewer_comment=comment, comment_basis=basis,
        )
        recovered.append({
            "review_group_id": group_id,
            "member_count": member_count,
            "raw_group_label": raw_label,
            "normalized_group_label": normalized,
            "confidence": confidence,
            "reason_code": reason,
            "reviewer_comment": comment,
            "comment_basis": basis,
            "partition_detail_status": partition_detail_status(
                normalized, member_count, partitions[group_id]
            ),
            "reference_type": REFERENCE_TYPE,
            "submission_sha256": submission_hash,
            "dataset_version": DATASET_VERSION,
            "normalization_reason": normalization_reason,
        })
    return tuple(sorted(recovered, key=lambda row: row["review_group_id"])), {
        "expected_id_count": 48,
        "unique_id_count": len(seen),
        "missing_id_count": 0,
        "foreign_id_count": 0,
        "duplicate_id_count": 0,
        "member_source_integrity": "PASS",
        "reviewer_visible_schema_integrity": "PASS",
        "authoritative_human_formula_count": 0,
        "pair_or_llm_leakage": 0,
    }


def load_development_mapping(mapping_path: Path, expected_ids):
    assert_holdout_firewall(mapping_path)
    with mapping_path.open(encoding="utf-8", newline="") as handle:
        development_rows = [
            row for row in csv.DictReader(handle)
            if row["development_or_holdout"] == "DEVELOPMENT"
        ]
    rows = {row["review_group_id"]: row for row in development_rows}
    if (
        len(development_rows) != 48 or len(rows) != 48
        or set(rows) != set(expected_ids)
    ):
        raise DevelopmentReferenceError("R18G_B_DEVELOPMENT_REFERENCE_INVALID: mapping join")
    return rows


def _write_csv(path, fields, rows):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def _matrix(rows, field):
    result = defaultdict(Counter)
    for row in rows:
        result[str(row[field])][row["normalized_group_label"]] += 1
    return {
        key: {
            label: values.get(label, 0)
            for label in LABELS
        }
        for key, values in sorted(result.items())
    }


def _distribution(rows, field, *, label_field="normalized_group_label"):
    result = defaultdict(Counter)
    for row in rows:
        result[row[label_field]][str(row[field])] += 1
    return {key: dict(sorted(value.items())) for key, value in sorted(result.items())}


def _source_bucket(row):
    values = {
        "ACCEPTED_LIKELY_GROUP": "current accepted likely",
        "ACCEPTED_REVIEW_GROUP": "current accepted review",
        "DEFERRED_FAMILY_CANDIDATE": "deferred-family hypothesis",
        "CONFLICT_CONTEXT_CANDIDATE": "conflict-context hypothesis",
        "ADR_PARTITION_CHALLENGER": "ADR-aligned shadow",
        "GENERIC_BICYCLE_FAMILY": "other bounded source",
    }
    if row["partition_sensitive"] == "true" and row["current_or_shadow"] == "CURRENT":
        return "current partition-sensitive"
    return values.get(row["source_kind"], "other bounded source")


def _support_comparison(shadow, current):
    if not current:
        return "COMPARISON_NOT_RESOLVABLE"
    labels = {row["normalized_group_label"] for row in current}
    if shadow["normalized_group_label"] == "INSUFFICIENT_INFORMATION" or (
        "INSUFFICIENT_INFORMATION" in labels
    ):
        return "HUMAN_INSUFFICIENT"
    shadow_supported = shadow["normalized_group_label"] == "SAME_ONE_IDENTITY"
    current_supported = "SAME_ONE_IDENTITY" in labels
    if shadow_supported and current_supported:
        return "BOTH_SUPPORTED"
    if shadow_supported:
        return "SHADOW_ONLY_HUMAN_SUPPORTED"
    if current_supported:
        return "CURRENT_ONLY_HUMAN_SUPPORTED"
    return "NEITHER_SUPPORTED"


def partition_comparisons(rows):
    comparisons = []
    current = [row for row in rows if row["current_or_shadow"] == "CURRENT"]
    for shadow in (
        row for row in rows
        if row["partition_sensitive"] == "true"
        and row["current_or_shadow"] == "SHADOW_ADR_ORDER"
    ):
        shadow_members = set(shadow["source_member_refs"].split("|"))
        competitors = [
            row for row in current
            if shadow_members & set(row["source_member_refs"].split("|"))
        ]
        comparisons.append({
            "comparison_work_unit_id": shadow["work_unit_id"],
            "shadow_review_group_id": shadow["review_group_id"],
            "shadow_label": shadow["normalized_group_label"],
            "shadow_confidence": shadow["confidence"],
            "shadow_member_count": shadow["member_count"],
            "shadow_member_refs": shadow["source_member_refs"],
            "shadow_comment_basis": shadow["comment_basis"],
            "current_review_group_ids": "|".join(row["review_group_id"] for row in competitors),
            "current_labels": "|".join(row["normalized_group_label"] for row in competitors),
            "current_confidences": "|".join(row["confidence"] for row in competitors),
            "current_member_sets": "||".join(row["source_member_refs"] for row in competitors),
            "current_comment_bases": "|".join(row["comment_basis"] for row in competitors),
            "comparison_outcome": _support_comparison(shadow, competitors),
        })
    return tuple(comparisons)


def partition_policy_decision(comparisons, targeted_followup):
    if targeted_followup:
        return (
            "R18G_B_PARTITION_DETAIL_REQUIRED",
            "R18G_B1_TARGETED_PARTITION_COMPLETION",
        )
    outcomes = Counter(row["comparison_outcome"] for row in comparisons)
    if len(comparisons) < 2:
        return "R18G_B_GROUP_EVIDENCE_INSUFFICIENT", "ARCHITECT_REVIEW"
    shadow = outcomes["SHADOW_ONLY_HUMAN_SUPPORTED"]
    current = outcomes["CURRENT_ONLY_HUMAN_SUPPORTED"]
    if shadow >= 2 and shadow > current and not current:
        return (
            "R18G_B_ADR_PARTITION_OBJECTIVE_SUPPORTED",
            "R18H_GF5_PARTITION_OBJECTIVE_CORRECTION",
        )
    if current >= 2 and current > shadow:
        return (
            "R18G_B_CURRENT_PARTITION_OBJECTIVE_SUPPORTED",
            "FREEZE_GF5_PARTITION_POLICY_THEN_HOLDOUT",
        )
    return (
        "R18G_B_PARTITION_CHANGE_NOT_JUSTIFIED",
        "FREEZE_GF5_PARTITION_POLICY_THEN_HOLDOUT",
    )


def evaluate_development(submission_path: Path, blank_path: Path, mapping_path: Path,
                         manifest_path: Path, output_dir: Path):
    assert_holdout_firewall(submission_path, blank_path, mapping_path, manifest_path)
    reference, integrity = validate_and_recover(submission_path, blank_path)
    mapping = load_development_mapping(
        mapping_path, (row["review_group_id"] for row in reference)
    )
    joined = []
    for row in reference:
        mapped = mapping[row["review_group_id"]]
        joined.append({**row, **{field: mapped[field] for field in MAPPING_FIELDS}})
    for row in joined:
        row["source_bucket"] = _source_bucket(row)
        row["cross_site_state"] = (
            "CROSS_SITE" if int(row["site_count"]) > 1 else "SAME_SITE"
        )
        row["review_support_origin"] = (
            "R18C_DEMOTED_PRESENT"
            if int(row["r18c_demoted_review_edge_count"]) > 0
            else "NATIVE_OR_NO_REVIEW"
        )

    comparisons = partition_comparisons(joined)
    targeted = tuple(
        row for row in joined
        if row["partition_sensitive"] == "true"
        and row["normalized_group_label"] == "NOT_ONE_IDENTITY"
        and int(row["member_count"]) >= 3
        and row["partition_detail_status"] == "UNAVAILABLE"
    )
    decision, next_step = partition_policy_decision(comparisons, targeted)
    accepted = [row for row in joined if row["source_kind"].startswith("ACCEPTED_")]
    accepted_by_status = {}
    for status in ("LIKELY_DUPLICATE_GROUP", "POSSIBLE_DUPLICATE_GROUP_REVIEW"):
        values = [row for row in accepted if row["system_status"] == status]
        counts = Counter(row["normalized_group_label"] for row in values)
        accepted_by_status[status] = {
            "total": len(values),
            "counts": {label: counts.get(label, 0) for label in LABELS},
            "rates": {
                label: round(counts.get(label, 0) / len(values), 4) if values else None
                for label in LABELS
            },
            "high_confidence_not_one": sum(
                row["normalized_group_label"] == "NOT_ONE_IDENTITY"
                and row["confidence"] == "HIGH" for row in values
            ),
            "comment_basis": dict(sorted(Counter(row["comment_basis"] for row in values).items())),
        }
    bicycle = next(
        row for row in joined if row["source_kind"] == "GENERIC_BICYCLE_FAMILY"
    )
    raw_counts = Counter(row["raw_group_label"] for row in reference)
    normalized_counts = Counter(row["normalized_group_label"] for row in reference)
    output_dir.mkdir(parents=True, exist_ok=True)
    reference_path = output_dir / "r18g_development_group_reference.csv"
    evaluation_path = output_dir / "r18g_development_group_evaluation.csv"
    comparison_path = output_dir / "r18g_partition_policy_comparison.csv"
    summary_path = output_dir / "r18g_development_summary.json"
    _write_csv(reference_path, REFERENCE_FIELDS, reference)
    evaluation_fields = (*REFERENCE_FIELDS, *MAPPING_FIELDS, "source_bucket",
                         "cross_site_state", "review_support_origin")
    _write_csv(evaluation_path, evaluation_fields, joined)
    comparison_fields = tuple(comparisons[0].keys()) if comparisons else (
        "comparison_work_unit_id", "shadow_review_group_id", "comparison_outcome",
    )
    _write_csv(comparison_path, comparison_fields, comparisons)

    summary = {
        "classification": decision,
        "next": next_step,
        "dataset_version": DATASET_VERSION,
        "reference_type": REFERENCE_TYPE,
        "claim_limit": [
            "DEVELOPMENT_ARCHITECTURE_EVIDENCE", "NOT_INDEPENDENT_VALIDATION",
            "NOT_DUAL_REVIEWED", "NOT_ADJUDICATED", "NOT_GOLD_STANDARD",
            "NOT_PRODUCTION_ACCURACY_TRUTH",
        ],
        "holdout_firewall": "SEALED_HUMAN_LABELS_UNSEEN",
        "provider_calls": 0,
        "runtime_or_gf5_changes": 0,
        "submission_sha256": sha256_file(submission_path),
        "blank_template_sha256": sha256_file(blank_path),
        "source_manifest_sha256": sha256_file(manifest_path),
        "integrity": integrity,
        "raw_label_distribution": dict(sorted(raw_counts.items())),
        "mixed_label_recovery_count": sum(
            bool(row["normalization_reason"]) for row in reference
        ),
        "normalized_label_distribution": {
            label: normalized_counts.get(label, 0) for label in LABELS
        },
        "confidence_distribution": dict(sorted(Counter(row["confidence"] for row in reference).items())),
        "comment_basis_distribution": dict(sorted(Counter(row["comment_basis"] for row in reference).items())),
        "partition_detail_distribution": dict(sorted(Counter(row["partition_detail_status"] for row in reference).items())),
        "whole_group_matrix_by_source": _matrix(joined, "source_bucket"),
        "whole_group_matrix_by_size": _matrix(joined, "member_count"),
        "whole_group_matrix_by_site_count": _matrix(joined, "site_count"),
        "whole_group_matrix_by_cross_site": _matrix(joined, "cross_site_state"),
        "whole_group_matrix_by_generic_state": _matrix(joined, "generic_only_state"),
        "whole_group_matrix_by_evidence_provenance": _matrix(joined, "evidence_provenance"),
        "whole_group_matrix_by_review_origin": _matrix(joined, "review_support_origin"),
        "whole_group_matrix_by_confidence": _matrix(joined, "confidence"),
        "whole_group_matrix_by_comment_basis": _matrix(joined, "comment_basis"),
        "current_accepted_quality": accepted_by_status,
        "deferred_family_matrix": _matrix(
            [row for row in joined if row["source_kind"] in {
                "DEFERRED_FAMILY_CANDIDATE", "GENERIC_BICYCLE_FAMILY"
            }], "source_kind"
        ),
        "bicycle_finding": {
            key: bicycle[key] for key in (
                "review_group_id", "raw_group_label", "normalized_group_label",
                "confidence", "reason_code", "reviewer_comment", "comment_basis",
                "current_or_shadow", "system_status", "partition_detail_status",
            )
        },
        "partition_comparison_outcomes": dict(sorted(Counter(
            row["comparison_outcome"] for row in comparisons
        ).items())),
        "partition_comparison_count": len(comparisons),
        "partition_comparison_preference_order": [
            "SAME_ONE_IDENTITY", "INSUFFICIENT_INFORMATION", "NOT_ONE_IDENTITY"
        ],
        "partition_shadow_cannot_link_violations": sum(
            row["current_or_shadow"] == "SHADOW_ADR_ORDER"
            and int(row["cannot_link_count"]) > 0 for row in joined
        ),
        "partition_detail_required": "YES",
        "targeted_partition_followup": [
            {"review_group_id": row["review_group_id"], "member_count": int(row["member_count"])}
            for row in targeted
        ],
        "reason_distribution_by_normalized_label": _distribution(joined, "reason_code"),
        "comment_basis_by_normalized_label": _distribution(joined, "comment_basis"),
        "high_confidence_by_normalized_label": {
            label: sum(
                row["normalized_group_label"] == label and row["confidence"] == "HIGH"
                for row in joined
            ) for label in LABELS
        },
        "domain_external_case_count": sum(row["comment_basis"] == "DOMAIN_EXTERNAL" for row in joined),
        "mixed_knowledge_case_count": sum(row["comment_basis"] == "MIXED" for row in joined),
        "unclear_basis_case_count": sum(row["comment_basis"] == "UNCLEAR" for row in joined),
    }
    summary["artifacts"] = {
        path.name: {"bytes": path.stat().st_size, "sha256": sha256_file(path)}
        for path in (reference_path, evaluation_path, comparison_path)
    }
    write_manifest(summary_path, summary)
    return {**summary, "summary_artifact": {
        "path": summary_path.name, "bytes": summary_path.stat().st_size,
        "sha256": sha256_file(summary_path),
    }}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--submission", required=True, type=Path)
    parser.add_argument("--blank-template", required=True, type=Path)
    parser.add_argument("--mapping", required=True, type=Path)
    parser.add_argument("--manifest", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    result = evaluate_development(
        args.submission, args.blank_template, args.mapping, args.manifest,
        args.output_dir,
    )
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
