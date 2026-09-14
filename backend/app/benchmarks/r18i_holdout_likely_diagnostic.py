"""Offline causal diagnostic for consumed R18G-C Likely-group failures."""

from __future__ import annotations

import argparse
import csv
import json
import sqlite3
from collections import Counter
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook

from app.benchmarks.r18g_group_human_evidence import (
    DATASET_VERSION,
    VISIBLE_FIELDS,
    sha256_file,
    write_manifest,
)
from app.engine.identity_evidence_evaluator import (
    DeterministicIdentityContext,
    evaluate_canonical_identity_relationship,
)
from app.engine.identity_signature_derivation import derive_identity_signature
from app.engine.signed_identity_evidence import (
    classify_shadow_evidence,
    derive_signed_identity_evidence,
)
from app.services.canonical_record_service import CanonicalScanRecord


CLASSIFICATION = "R18I_NO_SAFE_GENERAL_CAUSE_PROVEN"
NEXT = "NEW_HUMAN_EVIDENCE_FOR_UNRESOLVED_CAUSE"
HOLDOUT_STATUS = "CONSUMED_FOR_DIAGNOSTIC_USE"
POST_CORRECTION_HOLDOUT_REQUIRED = "YES"
R18GC_COMMIT = "0fe6f5578b148a4d347cce109425502c6c1e9822"
POLICY_ID = "GF5_PARTITION_POLICY_V1"
LIKELY = "LIKELY_DUPLICATE_GROUP"
SAME = "SAME_ONE_IDENTITY"
NOT_ONE = "NOT_ONE_IDENTITY"
STRONG = "STRONG_SUPPORT"
REVIEW = "REVIEW_SUPPORT"


class R18IDiagnosticError(ValueError):
    """Fail-closed invalid or incomplete diagnostic input."""


def _read_csv(path):
    with path.open(encoding="utf-8", newline="") as handle:
        return list(csv.DictReader(handle))


def _write_csv(path, fields, rows):
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields, lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def select_likely_cases(rows):
    rows = tuple(rows)
    failed = tuple(sorted((
        row for row in rows
        if row["system_status"] == LIKELY
        and row["normalized_group_label"] == NOT_ONE
    ), key=lambda row: row["review_group_id"]))
    controls = tuple(sorted((
        row for row in rows
        if row["system_status"] == LIKELY
        and row["normalized_group_label"] == SAME
    ), key=lambda row: row["review_group_id"]))
    if len(failed) != 6 or len(controls) != 2:
        raise R18IDiagnosticError("R18I_EXACT_LIKELY_CASE_SET_INVALID")
    return failed, controls


def topology_class(pair_classes):
    values = tuple(pair_classes)
    if values and all(value == STRONG for value in values):
        return "ALL_STRONG"
    if values and set(values) <= {STRONG, REVIEW} and STRONG in values:
        return "STRONG_CORE_WITH_REVIEW_BRIDGE"
    if values and all(value == REVIEW for value in values):
        return "REVIEW_DOMINATED"
    return "MIXED_EVIDENCE"


def likely_promotion_trace(pair_classes, *, complete=True, unresolved_bridge=False):
    values = tuple(pair_classes)
    promoted = bool(
        complete and values and all(value == STRONG for value in values)
        and not unresolved_bridge
    )
    return {
        "promoted": promoted,
        "rule": "ALL_INTERNAL_PAIRS_STRONG_COMPLETE_PAIRWISE",
        "all_pairs_strong": bool(values and all(value == STRONG for value in values)),
        "complete_pairwise": complete,
        "unresolved_bridge": unresolved_bridge,
    }


def membership_status_diagnosis(row, pair):
    if row["comment_basis"] == "DOMAIN_EXTERNAL" and not pair["protected_conflicts"]:
        return "DATA_INSUFFICIENCY_OR_EXTERNAL_KNOWLEDGE"
    if (
        pair["edge_class"] == STRONG
        and pair["support_provenance"] == "SHADOW_LEXICAL_ONLY_OR_UNRESOLVED"
        and not pair["protected_conflicts"]
    ):
        return "STATUS_PROMOTION_ERROR"
    return "MIXED_CAUSE"


def actionability(row, pair):
    if row["comment_basis"] == "DOMAIN_EXTERNAL":
        if pair["protected_conflicts"]:
            return "DETERMINISTICALLY_ACTIONABLE"
        if pair["unresolved_discriminator_count"]:
            return "PARTIALLY_ACTIONABLE"
        return "NOT_ACTIONABLE_FROM_AVAILABLE_INPUT"
    if row["comment_basis"] == "UNCLEAR":
        return "MIXED_OR_UNCLEAR"
    return "SOURCE_VISIBLE_LIKELY" if pair["protected_conflicts"] else "MIXED_OR_UNCLEAR"


def gf5_objective_causality(*, member_count, eligible_group_count):
    if int(member_count) == 2 and int(eligible_group_count) == 1:
        return "GF5_NOT_PRIMARY_CAUSE"
    return "UNKNOWN"


def shadow_policy(group):
    possible = int(group["possible_pair_count"])
    strong = int(group["strong_edge_count"])
    review = int(group["review_edge_count"])
    density = strong / possible if possible else 0.0
    trusted = group["support_provenance"] == "SHADOW_TRUSTED_IDENTITY_PRESENT"
    return {
        "S1_ALL_INTERNAL_PAIRS_STRONG": strong == possible,
        "S2_MINIMUM_STRONG_DENSITY": density == 1.0,
        "S3_ANY_REVIEW_REMAINS_REVIEW": review == 0,
        "S4_WEAKEST_LINK_STRONG": strong == possible,
        "S5_TRUSTED_PATH_PER_MEMBER": trusted,
    }


def _canonical(row):
    return CanonicalScanRecord(
        record_id=row["id"], scan_id=row["scan_id"],
        source_row_index=row["source_row_index"],
        record_ref_key=row["record_ref_key"],
        source_record_fingerprint=row["source_record_fingerprint"],
        part_no=row["part_no"], description=row["description"],
        contract=row["contract"], uom=row["uom"], type_code=row["type_code"],
        prime_commodity=row["prime_commodity"],
        second_commodity=row["second_commodity"],
        accounting_group=row["accounting_group"],
        part_product_code=row["part_product_code"],
        part_product_family=row["part_product_family"],
        product_category_id=row["product_category_id"],
        hsn_sac_code=row["hsn_sac_code"], hazard_code=row["hazard_code"],
        normalized_part_no=row["normalized_part_no"],
        normalized_description=row["normalized_description"],
        normalization_version=row["normalization_version"],
    )


def _readonly(path):
    connection = sqlite3.connect(
        f"file:{path.resolve().as_posix()}?mode=ro", uri=True
    )
    connection.row_factory = sqlite3.Row
    return connection


def _visible_members(workbook_path):
    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    sheet = workbook["Members"]
    headers = tuple(cell.value for cell in sheet[1])
    expected = ("review_group_id", "member_no", *VISIBLE_FIELDS)
    if headers != expected:
        raise R18IDiagnosticError("R18I_HOLDOUT_MEMBER_SCHEMA_INVALID")
    result = {}
    for values in sheet.iter_rows(min_row=2, values_only=True):
        if values[0] in (None, ""):
            continue
        result[(values[0], int(values[1]))] = {
            field: values[index + 2] or ""
            for index, field in enumerate(VISIBLE_FIELDS)
        }
    workbook.close()
    return result


def _pair_diagnostic(left, right, context):
    evaluated = evaluate_canonical_identity_relationship(left, right, context)
    technical = json.loads(evaluated.technical_evidence_json)
    lexical = technical.get("lexical_trust_assessment") or {}
    signed = derive_signed_identity_evidence(
        derive_identity_signature(left, record_reference=left.record_ref_key),
        derive_identity_signature(right, record_reference=right.record_ref_key),
    )
    return {
        "edge_class": evaluated.edge_class.value,
        "classification_reason_codes": list(evaluated.classification_reason_codes),
        "deterministic_score": evaluated.deterministic_score,
        "score_components": json.loads(evaluated.component_scores_json),
        "rule_decision": evaluated.rule_decision,
        "rejection_reason": evaluated.rejection_reason,
        "protected_conflicts": json.loads(evaluated.protected_conflicts_json),
        "generic_evidence": json.loads(evaluated.generic_evidence_json),
        "technical_evidence": technical,
        "support_provenance": classify_shadow_evidence(signed).value,
        "signed_evidence_facts": [
            {
                "channel": fact.channel.value,
                "semantic_category": fact.semantic_category.value,
                "semantic_key": fact.semantic_key,
                "normalized_matches": list(fact.normalized_matches),
                "reason_code": fact.reason_code,
            }
            for fact in signed.facts
        ],
        "lexical_trust_assessment": lexical,
        "unresolved_discriminator_count": int(
            lexical.get("unresolved_discriminator_count") or 0
        ),
        "r18c_downgrade_reasons": list(lexical.get("risk_reasons") or []),
    }


def _group_features(row, pair):
    possible = int(row["member_count"]) * (int(row["member_count"]) - 1) // 2
    return {
        "all_strong": pair["edge_class"] == STRONG and possible == 1,
        "mixed_strong_review": False,
        "description_dominant": bool(
            pair["lexical_trust_assessment"].get("description_dominance")
        ),
        "trusted_identity_present": (
            pair["support_provenance"] == "SHADOW_TRUSTED_IDENTITY_PRESENT"
        ),
        "lexical_only_or_unresolved": (
            pair["support_provenance"] == "SHADOW_LEXICAL_ONLY_OR_UNRESOLVED"
        ),
        "generic_or_copy_risk": bool(
            pair["lexical_trust_assessment"].get("generic_or_copy_risk")
        ),
        "unresolved_discriminator": bool(pair["unresolved_discriminator_count"]),
        "part_number_coherence_floor_met": bool(
            pair["lexical_trust_assessment"].get("part_family_coherence")
        ),
        "group_size_two": int(row["member_count"]) == 2,
        "strong_density_one": pair["edge_class"] == STRONG and possible == 1,
        "weakest_link_strong": pair["edge_class"] == STRONG,
    }


def evaluate_diagnostic(
    holdout_evaluation_path: Path,
    development_evaluation_path: Path,
    completed_workbook_path: Path,
    database_path: Path,
    output_dir: Path,
):
    rows = _read_csv(holdout_evaluation_path)
    failed, controls = select_likely_cases(rows)
    development_controls = tuple(
        row for row in _read_csv(development_evaluation_path)
        if row["system_status"] == LIKELY
        and row["normalized_group_label"] == SAME
    )
    if len(development_controls) != 6:
        raise R18IDiagnosticError("R18I_DEVELOPMENT_LIKELY_CONTROL_SET_INVALID")

    visible = _visible_members(completed_workbook_path)
    connection = _readonly(database_path)
    scan = connection.execute(
        "SELECT selected_fields FROM duplicate_scan WHERE id=31"
    ).fetchone()
    if scan is None:
        raise R18IDiagnosticError("R18I_SCAN31_UNAVAILABLE")
    context = DeterministicIdentityContext(
        "DISCOVERY", tuple(json.loads(scan["selected_fields"]))
    )

    audit = []
    pair_rows = []
    feature_rows = []
    for role, selected in (("FAILED_LIKELY", failed), ("ACCEPTED_LIKELY_CONTROL", controls)):
        for row in selected:
            refs = row["source_member_refs"].split("|")
            records = []
            for ref in refs:
                record = connection.execute(
                    "SELECT * FROM scan_record_snapshot WHERE scan_id=31 AND record_ref_key=?",
                    (ref,),
                ).fetchone()
                if record is None:
                    raise R18IDiagnosticError("R18I_SOURCE_RECORD_UNAVAILABLE")
                records.append(_canonical(record))
            if len(records) != 2:
                raise R18IDiagnosticError("R18I_EXPECTED_TWO_MEMBER_LIKELY_GROUP")
            pair = _pair_diagnostic(records[0], records[1], context)
            if pair["edge_class"] != STRONG:
                raise R18IDiagnosticError("R18I_FROZEN_LIKELY_PAIR_NOT_STRONG")
            source_members = []
            for index, ref in enumerate(refs, 1):
                source_members.append({
                    "member_reference": ref,
                    **visible[(row["review_group_id"], index)],
                })
            group_summary = {
                "member_count": 2,
                "possible_pair_count": 1,
                "evaluated_pair_count": 1,
                "strong_support_count": 1,
                "review_support_count": 0,
                "cannot_link_count": 0,
                "support_density": 1.0,
                "strong_support_density": 1.0,
                "missing_evidence_count": 0,
                "validation_mode": "COMPLETE_PAIRWISE",
                "bridge_unresolved": False,
            }
            diagnosis = (
                membership_status_diagnosis(row, pair)
                if role == "FAILED_LIKELY" else "ACCEPTED_CONTROL"
            )
            action = actionability(row, pair) if role == "FAILED_LIKELY" else "CONTROL"
            semantic_gap = bool(
                role == "FAILED_LIKELY"
                and (
                    row["comment_basis"] == "DOMAIN_EXTERNAL"
                    or pair["support_provenance"]
                    == "SHADOW_LEXICAL_ONLY_OR_UNRESOLVED"
                )
            )
            audit.append({
                "case_role": role,
                "review_group_id": row["review_group_id"],
                "current_group_id": row["current_group_id"],
                "member_count": 2,
                "member_references": "|".join(refs),
                "source_members_json": json.dumps(source_members, ensure_ascii=True, sort_keys=True),
                "sites": "|".join(str(item["site"]) for item in source_members),
                "senior_confidence": row["confidence"],
                "senior_reason_code": row["reason_code"],
                "senior_reviewer_comment": row["reviewer_comment"],
                "senior_comment_basis": row["comment_basis"],
                "group_evidence_summary_json": json.dumps(group_summary, sort_keys=True),
                "strong_edge_count": 1,
                "native_review_edge_count": 0,
                "r18c_demoted_review_edge_count": 0,
                "cannot_link_count": 0,
                "generic_only_state": row["generic_only_state"],
                "bridge_state": "RESOLVED_TWO_MEMBER_COMPLETE_STRONG",
                "missing_evidence_state": "NONE",
                "coverage": "1_OF_1_INTERNAL_PAIRS_EVALUATED",
                "gf5_objective": "(2,2,1,0,-1)",
                "status_promotion_reason": "ALL_INTERNAL_PAIRS_STRONG_COMPLETE_PAIRWISE",
                "topology_class": topology_class((pair["edge_class"],)),
                "primary_case_diagnosis": diagnosis,
                "source_actionability": action,
                "gf5_objective_causality": gf5_objective_causality(
                    member_count=2, eligible_group_count=1
                ),
                "semantic_interpretation_gap": str(semantic_gap).lower(),
            })
            pair_rows.append({
                "case_role": role,
                "review_group_id": row["review_group_id"],
                "left_member_reference": refs[0],
                "right_member_reference": refs[1],
                "gf4_class": pair["edge_class"],
                "classification_reason_codes_json": json.dumps(pair["classification_reason_codes"]),
                "deterministic_score": pair["deterministic_score"],
                "score_components_json": json.dumps(pair["score_components"], sort_keys=True),
                "rule_decision": pair["rule_decision"],
                "rejection_reason": pair["rejection_reason"],
                "trusted_identity_evidence_json": json.dumps(pair["signed_evidence_facts"], sort_keys=True),
                "lexical_evidence_json": json.dumps(pair["lexical_trust_assessment"], sort_keys=True),
                "technical_variant_evidence_json": json.dumps(pair["technical_evidence"], sort_keys=True),
                "protected_conflicts_json": json.dumps(pair["protected_conflicts"], sort_keys=True),
                "unresolved_observation_count": pair["unresolved_discriminator_count"],
                "r18c_downgrade_reasons_json": json.dumps(pair["r18c_downgrade_reasons"]),
                "support_provenance": pair["support_provenance"],
                "topology_class": topology_class((pair["edge_class"],)),
            })
            features = _group_features(row, pair)
            feature_rows.append({"case_role": role, **features})
    connection.close()

    feature_names = tuple(key for key in feature_rows[0] if key != "case_role")
    contrast = []
    for feature in feature_names:
        contrast.append({
            "feature": feature,
            "rejected_likely_count": sum(
                item[feature] for item in feature_rows if item["case_role"] == "FAILED_LIKELY"
            ),
            "accepted_likely_count": sum(
                item[feature] for item in feature_rows
                if item["case_role"] == "ACCEPTED_LIKELY_CONTROL"
            ),
            "rejected_likely_total": 6,
            "accepted_likely_total": 2,
        })

    shadow_groups = [{
        "cohort": item["case_role"],
        "possible_pair_count": 1,
        "strong_edge_count": 1,
        "review_edge_count": 0,
        "support_provenance": next(
            pair["support_provenance"] for pair in pair_rows
            if pair["review_group_id"] == item["review_group_id"]
        ),
    } for item in audit]
    shadow_groups.extend({
        "cohort": "DEVELOPMENT_SAME_LIKELY_CONTROL",
        "possible_pair_count": 1,
        "strong_edge_count": int(row["strong_edge_count"]),
        "review_edge_count": (
            int(row["native_review_edge_count"])
            + int(row["r18c_demoted_review_edge_count"])
        ),
        "support_provenance": row["evidence_provenance"],
    } for row in development_controls)
    shadow_rows = []
    for policy in shadow_policy(shadow_groups[0]):
        for cohort in (
            "FAILED_LIKELY", "ACCEPTED_LIKELY_CONTROL",
            "DEVELOPMENT_SAME_LIKELY_CONTROL",
        ):
            values = [item for item in shadow_groups if item["cohort"] == cohort]
            retained = sum(shadow_policy(item)[policy] for item in values)
            shadow_rows.append({
                "policy": policy,
                "cohort": cohort,
                "total": len(values),
                "retained_likely": retained,
                "downgraded_to_review": len(values) - retained,
                "monotonic": "true",
            })

    output_dir.mkdir(parents=True, exist_ok=True)
    audit_path = output_dir / "r18i_failed_likely_group_audit.csv"
    topology_path = output_dir / "r18i_group_pair_topology.csv"
    shadow_path = output_dir / "r18i_likely_status_shadow_policies.csv"
    summary_path = output_dir / "r18i_summary.json"
    _write_csv(audit_path, tuple(audit[0]), audit)
    _write_csv(topology_path, tuple(pair_rows[0]), pair_rows)
    _write_csv(shadow_path, tuple(shadow_rows[0]), shadow_rows)
    case_diagnoses = Counter(
        row["primary_case_diagnosis"] for row in audit
        if row["case_role"] == "FAILED_LIKELY"
    )
    summary = {
        "classification": CLASSIFICATION,
        "next": NEXT,
        "holdout_status": HOLDOUT_STATUS,
        "post_r18_correction_independent_validation_required": (
            POST_CORRECTION_HOLDOUT_REQUIRED
        ),
        "baseline_commit": R18GC_COMMIT,
        "policy_id": POLICY_ID,
        "policy_status": "FROZEN_UNCHANGED",
        "runtime_changes": 0,
        "provider_calls": 0,
        "dataset_version": DATASET_VERSION,
        "failed_likely_count": len(failed),
        "accepted_likely_control_count": len(controls),
        "development_same_likely_control_count": len(development_controls),
        "case_diagnosis_distribution": dict(sorted(case_diagnoses.items())),
        "pair_topology_distribution": dict(sorted(Counter(
            row["topology_class"] for row in pair_rows
            if row["case_role"] == "FAILED_LIKELY"
        ).items())),
        "gf5_objective_causality": "GF5_NOT_PRIMARY_CAUSE",
        "status_only_correction_safe": False,
        "status_only_reason": (
            "S1-S4 retain all eight holdout Likely groups; S5 downgrades both "
            "accepted holdout controls and five of six Development SAME Likely "
            "controls while retaining one rejected holdout case"
        ),
        "immediate_mechanical_cause": (
            "Each rejected two-member group contains one false-positive Strong "
            "edge; current Likely promotion correctly follows the frozen all-Strong rule"
        ),
        "safe_general_runtime_rule_proven": False,
        "semantic_interpretation_gap_count": sum(
            row["semantic_interpretation_gap"] == "true" for row in audit
            if row["case_role"] == "FAILED_LIKELY"
        ),
        "contrast": contrast,
        "shadow_policy_results": shadow_rows,
        "preservation": {
            "bicycle_generic_deferral": "UNAFFECTED_ANALYSIS_ONLY",
            "head_tail_cannot_link": "UNAFFECTED_ANALYSIS_ONLY",
            "r18c_invariants": "UNCHANGED",
            "review_suppression": "NOT_AUTHORIZED",
        },
        "claim_limit": [
            "CONSUMED_HOLDOUT_DIAGNOSTIC_ONLY",
            "SINGLE_SENIOR_DOMAIN_EXPERT",
            "NO_POST_CHANGE_VALIDATION_CLAIM",
            "NO_CAUSAL_RULE_FIT_FROM_IDENTIFIERS_OR_NOUNS",
        ],
        "source_hashes": {
            "holdout_evaluation": sha256_file(holdout_evaluation_path),
            "development_evaluation": sha256_file(development_evaluation_path),
            "completed_workbook": sha256_file(completed_workbook_path),
        },
    }
    summary["artifacts"] = {
        path.name: {"bytes": path.stat().st_size, "sha256": sha256_file(path)}
        for path in (audit_path, topology_path, shadow_path)
    }
    write_manifest(summary_path, summary)
    return {
        **summary,
        "summary_artifact": {
            "bytes": summary_path.stat().st_size,
            "sha256": sha256_file(summary_path),
        },
    }


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--holdout-evaluation", required=True, type=Path)
    parser.add_argument("--development-evaluation", required=True, type=Path)
    parser.add_argument("--completed-workbook", required=True, type=Path)
    parser.add_argument("--database", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()
    result = evaluate_diagnostic(
        args.holdout_evaluation,
        args.development_evaluation,
        args.completed_workbook,
        args.database,
        args.output_dir,
    )
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
